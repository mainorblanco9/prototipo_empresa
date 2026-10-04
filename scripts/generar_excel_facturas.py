"""Genera los Excel de prueba de GuanaBill (todos los tipos de factura)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["factura", "emisor", "cedula_emisor", "receptor", "cedula_receptor",
           "fecha", "descripcion", "cantidad", "precio_con_iva", "iva"]
OUT = Path(__file__).resolve().parent.parent / "samples"

BANCO = ("Banco Nacional de Costa Rica", "4000001021")
BAC = ("BAC Credomatic S.A.", "3101021124")
ICE = ("Instituto Costarricense de Electricidad", "4000042139")
AYA = ("Acueductos y Alcantarillados", "4000042138")
COMERCIAL = ("Comercial Guanacaste S.A.", "3101123456")
CLINICA = ("Clínica Santa Rosa S.A.", "3101333444")
SUPER = ("Supermercado La Canasta S.A.", "3101555666")
COLEGIO = ("Colegio Bilingüe del Pacífico", "3101777888")
SEGUROS = ("Seguros del Istmo S.A.", "3101999000")
FISICA = ("Minor Andres Blanco Chaves", "114560789")
EXTRANJERO = ("Consultora Nómada", "155812345678")

CLI_J = ("Cliente de Prueba S.A.", "3101654321")
CLI_F = ("María Fernanda Rojas Vargas", "205670123")
CLI_D = ("Carlos Alberto Mendoza (DIMEX)", "155800112233")
CLI_E = ("Hotel Bahía Verde S.A.", "3101246810")

VALID = []


def inv(num, issuer, receiver, date, lines):
    for desc, qty, price, iva in lines:
        VALID.append([num, issuer[0], issuer[1], receiver[0], receiver[1], date, desc, qty, price, iva])


# Tarifas
inv("T-13-001", COMERCIAL, CLI_J, "2026-08-05", [("Servicio de consultoría", 1, 113000, "13")])
inv("T-04-002", CLINICA, CLI_F, "2026-08-06", [("Consulta médica especializada", 1, 41600, "4")])
inv("T-02-003", SEGUROS, CLI_J, "2026-08-07", [("Prima póliza de vehículo", 1, 204000, "2")])
inv("T-01-004", SUPER, CLI_F, "2026-08-08", [("Arroz 2 kg", 6, 2020, "1"), ("Frijoles 1 kg", 6, 1515, "1"), ("Leche 1 L", 12, 1010, "1")])
inv("T-00-005", COMERCIAL, CLI_J, "2026-08-09", [("Venta con tarifa 0%", 1, 50000, "0")])
inv("T-EX-006", COLEGIO, CLI_F, "2026-08-10", [("Mensualidad agosto", 1, 185000, "exempt"), ("Matrícula", 1, 60000, "exempt")])
# Multilínea y mixtas
inv("M-001", COMERCIAL, CLI_J, "2026-08-12", [("Licencia de software anual", 1, 565000, "13"), ("Capacitación de equipo", 2, 40000, "exempt"), ("Soporte técnico", 10, 11300, "13")])
inv("M-002", SUPER, CLI_E, "2026-08-14", [("Abarrotes", 20, 3030, "1"), ("Artículos de limpieza", 8, 4520, "13"), ("Medicamentos", 3, 5200, "2")])
# Identificaciones
inv("ID-F-001", FISICA, CLI_F, "2026-08-18", [("Servicios profesionales de programación", 1, 226000, "13")])
inv("ID-D-002", EXTRANJERO, CLI_D, "2026-08-19", [("Asesoría remota", 3, 56500, "13")])
inv("ID-J-003", COMERCIAL, CLI_D, "2026-08-20", [("Equipo de cómputo", 1, 452000, "13")])
# Bancos
inv("B-BN-001", BANCO, CLI_J, "2026-09-02", [("Comisión por transferencias SINPE", 120, 113, "13"), ("Mantenimiento de cuenta empresarial", 1, 8475, "13")])
inv("B-BN-002", BANCO, CLI_F, "2026-09-03", [("Comisión tarjeta de crédito", 1, 5650, "13"), ("Intereses (exentos)", 1, 42300, "exempt")])
inv("B-BAC-003", BAC, CLI_E, "2026-09-04", [("Administración de planilla", 1, 135600, "13"), ("Cobro de cheques", 15, 565, "13")])
inv("B-BAC-004", BAC, CLI_J, "2026-09-05", [("Servicio de custodia", 1, 79100, "13")])
# Servicios públicos
inv("U-ICE-001", ICE, CLI_J, "2026-09-08", [("Consumo eléctrico industrial", 1, 385400, "13"), ("Alumbrado público", 1, 12800, "exempt")])
inv("U-AYA-002", AYA, CLI_J, "2026-09-09", [("Consumo de agua", 1, 24500, "1")])
# Montos extremos y cantidades decimales
inv("X-BIG-001", COMERCIAL, CLI_E, "2026-09-12", [("Maquinaria industrial", 1, 38500000, "13")])
inv("X-DEC-002", SUPER, CLI_F, "2026-09-13", [("Carne de res (kg)", "2.35", 6850, "1"), ("Queso fresco (kg)", "0.75", 7400, "1")])
inv("X-SMALL-003", COMERCIAL, CLI_F, "2026-09-14", [("Cargo por servicio", 1, 113, "13")])
inv("X-MANY-004", SUPER, CLI_J, "2026-09-15", [(f"Producto {i:02d}", i, 1000 + i * 125, ["13", "1", "2", "4", "exempt"][i % 5]) for i in range(1, 13)])
# Octubre (periodo abierto)
inv("O-001", COMERCIAL, CLI_J, "2026-10-01", [("Desarrollo de módulo web", 40, 33900, "13")])
inv("O-002", CLINICA, CLI_F, "2026-10-02", [("Exámenes de laboratorio", 1, 62400, "4"), ("Consulta general", 1, 31200, "4")])
inv("O-003", COLEGIO, CLI_F, "2026-10-03", [("Mensualidad octubre", 1, 185000, "exempt")])
inv("O-004", SEGUROS, CLI_E, "2026-10-03", [("Seguro de incendio", 1, 306000, "2")])

# Casos con comportamiento especial
SPECIAL = []
SPECIAL_REJECT = [
    ["R-001", COMERCIAL[0], COMERCIAL[1], CLI_J[0], CLI_J[1], "2026-09-20", "Servicio RECHAZAR (el simulador lo rechaza)", 1, 56500, "13"],
]
INVALID = [
    ["E-001", COMERCIAL[0], COMERCIAL[1], CLI_J[0], CLI_J[1], "fecha-mala", "Fecha inválida", 1, 1000, "13"],
    ["", COMERCIAL[0], COMERCIAL[1], CLI_J[0], CLI_J[1], "2026-09-21", "Sin número de factura", 1, 1000, "13"],
    ["E-003", COMERCIAL[0], COMERCIAL[1], CLI_J[0], CLI_J[1], "2026-09-21", "Tarifa inexistente", 1, 1000, "15"],
    ["E-004", COMERCIAL[0], COMERCIAL[1], CLI_J[0], CLI_J[1], "2026-09-21", "Precio negativo", 1, -5000, "13"],
    ["E-005", COMERCIAL[0], COMERCIAL[1], CLI_J[0], CLI_J[1], "2026-09-21", "", 1, 1000, "13"],
    ["E-006", COMERCIAL[0], COMERCIAL[1], CLI_J[0], CLI_J[1], "2026-09-21", "Cantidad cero", 0, 1000, "13"],
]
# Importa bien pero falla la validación local: cédula del emisor inválida (como el 422 real)
BAD_ISSUER = [
    ["V-001", "Minor Andres Blanco Chaves", "1", CLI_J[0], CLI_J[1], "2026-09-22", "Emisor con cédula inválida", 1, 22600, "13"],
    ["V-002", COMERCIAL[0], COMERCIAL[1], "Cliente sin cédula", "12", "2026-09-22", "Receptor con cédula inválida", 1, 22600, "13"],
]


def write(path, rows, guide):
    wb = Workbook()
    ws = wb.active
    ws.title = "Facturas"
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    fill = PatternFill("solid", fgColor="0B6FB8")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
    ws.freeze_panes = "A2"
    widths = [12, 34, 14, 34, 15, 12, 46, 9, 15, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    g = wb.create_sheet("Guia")
    for line in guide:
        g.append([line])
    g.column_dimensions["A"].width = 120
    g["A1"].font = Font(bold=True, size=13)
    for row in g.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


GUIDE_ALL = [
    "GuanaBill · facturas de prueba de todo tipo",
    "Una fila por línea; las filas con el mismo número de factura forman una sola factura.",
    "Precios con IVA incluido. Tarifas: 13, 4, 2, 1, 0 o exempt.",
    "Identificaciones: física 9 dígitos, jurídica 10, DIMEX/NITE 11-12.",
    "T-*: una factura por cada tarifa. M-*: líneas mixtas. ID-*: tipos de cédula. B-*: bancos. U-*: servicios públicos.",
    "X-*: montos extremos, decimales y muchas líneas. O-*: periodo de octubre.",
    "R-001: contiene la palabra RECHAZAR, el simulador local la rechaza (prueba el análisis de IA del rechazo).",
    "V-*: se importan, pero la validación local las bloquea (cédula inválida) antes de enviar.",
    "E-*: errores de formato; el importador las reporta y no las crea.",
    "Importar desde 'Importar Excel' (vista previa primero, luego confirmar).",
]
write(OUT / "guanabill_facturas_todo_tipo.xlsx", VALID + SPECIAL_REJECT + BAD_ISSUER + INVALID, GUIDE_ALL)
write(OUT / "guanabill_facturas_validas.xlsx", VALID, GUIDE_ALL[:6])
print(len({r[0] for r in VALID}), "facturas validas;", len(VALID), "lineas")
