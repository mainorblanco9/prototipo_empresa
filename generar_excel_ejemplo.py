from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


headers = [
    "fecha",
    "proveedor",
    "numero_factura",
    "moneda",
    "subtotal",
    "impuesto",
    "total",
]
rows_with_errors = [
    ["2026-09-01", "Papelería Central", "FAC-1001", "CRC", 1000, 160, 1160],
    ["02/09/2026", "Servicios del Norte", "FAC-2048", "CRC", 2500, 400, 2900],
    ["2026-09-03", "Café y Oficina", "FAC-3307", "CRC", 450, 72, 522],
    ["2026-09-04", "Papelería Central", "FAC-1001", "CRC", 1000, 160, 1160],
    ["fecha-invalida", "Transporte Ejemplo", "FAC-4110", "CRC", 800, 128, 950],
    ["2026-09-06", "", "FAC-5112", "CRC", 300, 48, 348],
    ["2026-09-07", "Hospedaje Demo", "FAC-6201", "CRC", -500, 0, -500],
]

valid_rows = [
    ["2026-09-01", "Papelería Central", "FAC-1001", "CRC", 1000, 160, 1160],
    ["02/09/2026", "Servicios del Norte", "FAC-2048", "CRC", 2500, 400, 2900],
    ["2026-09-03", "Café y Oficina", "FAC-3307", "CRC", 450, 72, 522],
]


def write_workbook(filename, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Gastos"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="3157D5")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill

    for index, width in enumerate([16, 26, 20, 12, 14, 14, 14], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(filename)
    print(f"Archivo creado: {filename}")


write_workbook("ejemplo_gastos.xlsx", rows_with_errors)
write_workbook("ejemplo_gastos_validos.xlsx", valid_rows)
