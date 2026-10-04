import csv
import io
import re
import unicodedata
import uuid
from collections import OrderedDict
from datetime import date, datetime

from django.db import IntegrityError, transaction
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response
from rest_framework.views import APIView

from . import periods
from .api import (
    LineSerializer,
    _build_document,
    _document_fields,
)
from .issued_invoice import MAX_INVOICE_LINES, IssuedInvoiceDraftError
from .models import IssuedInvoice, IssuedInvoiceStatusEvent

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 2000

TEMPLATE_HEADERS = [
    "factura", "emisor", "cedula_emisor", "receptor", "cedula_receptor",
    "fecha", "descripcion", "cantidad", "precio_con_iva", "iva",
]
TEMPLATE_ROWS = [
    ["F-001", "Comercial Guanacaste S.A.", "3101123456", "Cliente de Prueba S.A.", "3101654321",
     "2026-01-15", "Servicio de consultoría", "1", "113000", "13"],
    ["F-001", "Comercial Guanacaste S.A.", "3101123456", "Cliente de Prueba S.A.", "3101654321",
     "2026-01-15", "Capacitación", "2", "40000", "exempt"],
    ["F-002", "Comercial Guanacaste S.A.", "3101123456", "Otro Cliente S.A.", "3101777888",
     "2026-01-20", "Canasta básica", "10", "5000", "1"],
]

ALIASES = {
    "factura": {"factura", "numero", "numero_factura", "id", "referencia"},
    "emisor": {"emisor", "emisor_nombre", "proveedor"},
    "cedula_emisor": {"cedula_emisor", "emisor_id", "id_emisor", "cedula_proveedor"},
    "receptor": {"receptor", "receptor_nombre", "cliente"},
    "cedula_receptor": {"cedula_receptor", "receptor_id", "id_receptor", "cedula_cliente"},
    "fecha": {"fecha", "fecha_emision", "issue_date"},
    "descripcion": {"descripcion", "detalle", "concepto", "description"},
    "cantidad": {"cantidad", "qty", "quantity"},
    "precio_con_iva": {"precio_con_iva", "precio", "precio_unitario", "unit_price", "monto"},
    "iva": {"iva", "tarifa", "tarifa_iva", "tax_rate"},
}


def _norm(header):
    text = unicodedata.normalize("NFKD", str(header or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def _map_headers(headers):
    mapping = {}
    for index, header in enumerate(headers):
        name = _norm(header)
        for canonical, names in ALIASES.items():
            if name in names and canonical not in mapping:
                mapping[canonical] = index
    return mapping


def read_rows(upload):
    name = (upload.name or "").lower()
    raw = upload.read()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    text = raw.decode("utf-8-sig", errors="replace")
    delimiter = ";" if text.splitlines()[0].count(";") > text.splitlines()[0].count(",") else ","
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _cell(row, mapping, key):
    index = mapping.get(key)
    if index is None or index >= len(row):
        return ""
    value = row[index]
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return "" if value is None else str(value).strip()


def parse_upload(upload):
    """Devuelve (facturas, errores_globales). Cada factura trae sus propios errores."""
    rows = read_rows(upload)
    if not rows:
        return [], ["El archivo está vacío."]
    mapping = _map_headers(rows[0])
    missing = [h for h in TEMPLATE_HEADERS if h not in mapping]
    if missing:
        return [], [f"Faltan columnas: {', '.join(missing)}. Descarga la plantilla."]
    body = [r for r in rows[1:] if any(str(c or "").strip() for c in r)]
    if len(body) > MAX_ROWS:
        return [], [f"Máximo {MAX_ROWS} filas por archivo."]

    groups = OrderedDict()
    for number, row in enumerate(body, start=2):
        key = _cell(row, mapping, "factura")
        group = groups.setdefault(key, {"key": key, "rows": [], "errors": [], "lines": []})
        group["rows"].append(number)
        group["lines"].append({
            "description": _cell(row, mapping, "descripcion"),
            "quantity": _cell(row, mapping, "cantidad") or "1",
            "unit_price": _cell(row, mapping, "precio_con_iva").replace(",", "."),
            "tax_rate": (_cell(row, mapping, "iva") or "13").replace("%", ""),
        })
        group.setdefault("header", {
            field: _cell(row, mapping, source)
            for field, source in (
                ("issuer_name", "emisor"), ("issuer_id", "cedula_emisor"),
                ("receiver_name", "receptor"), ("receiver_id", "cedula_receptor"),
                ("issue_date", "fecha"),
            )
        })

    invoices = []
    for group in groups.values():
        errors = group["errors"]
        if not group["key"]:
            errors.append(f"Fila {group['rows'][0]}: falta el número de factura.")
        header = group["header"]
        try:
            header["issue_date"] = date.fromisoformat(header["issue_date"][:10])
        except ValueError:
            errors.append("Fecha inválida (usa AAAA-MM-DD).")
        lines = LineSerializer(data=group["lines"], many=True)
        if len(group["lines"]) > MAX_INVOICE_LINES:
            errors.append(f"Máximo {MAX_INVOICE_LINES} líneas por factura.")
        elif not lines.is_valid():
            errors.append(f"Líneas inválidas (filas {group['rows'][0]}-{group['rows'][-1]}): revisa cantidad, precio e IVA.")
        document = None
        if not errors:
            try:
                document = _build_document({**header, "lines": lines.validated_data})
            except IssuedInvoiceDraftError as exc:
                errors.append(str(exc))
        invoices.append({
            "key": group["key"], "header": header, "errors": errors, "document": document,
            "lines": len(group["lines"]),
        })
    return invoices, []


class InvoiceImportView(APIView):
    """Importa Excel/CSV: una fila por línea, agrupadas por número de factura."""

    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        upload = request.FILES.get("file")
        if not upload:
            return Response({"detail": "Adjunta un archivo CSV o XLSX en 'file'."}, status=400)
        if upload.size > MAX_FILE_BYTES:
            return Response({"detail": "El archivo supera 2 MB."}, status=400)
        if not (upload.name or "").lower().endswith((".csv", ".xlsx")):
            return Response({"detail": "Solo se aceptan archivos .csv o .xlsx."}, status=400)
        try:
            invoices, global_errors = parse_upload(upload)
        except Exception:
            return Response({"detail": "No se pudo leer el archivo."}, status=400)
        if global_errors:
            return Response({"detail": " ".join(global_errors)}, status=400)

        confirm = str(request.data.get("confirm", "")).lower() in {"1", "true", "yes"}
        results, created_count = [], 0
        for item in invoices:
            row = {"factura": item["key"], "lines": item["lines"], "errors": list(item["errors"]), "status": "valid"}
            document = item["document"]
            if document:
                row["total"] = document["totals"]["total"]
                row["receiver"] = document["receiver"]["name"]
                if periods.is_closed(request.user, periods.period_of(item["header"]["issue_date"])):
                    row["errors"].append("El periodo de esa fecha está cerrado.")
            if row["errors"]:
                row["status"] = "error"
            elif confirm:
                client_id = uuid.uuid5(uuid.NAMESPACE_URL, f"guanabill:{request.user.pk}:{item['key']}")
                try:
                    with transaction.atomic():
                        invoice, created = IssuedInvoice.objects.get_or_create(
                            owner=request.user, client_id=client_id, defaults=_document_fields(document)
                        )
                        if created:
                            IssuedInvoiceStatusEvent.objects.create(
                                invoice=invoice, from_status="", to_status=invoice.status, actor=request.user
                            )
                except IntegrityError:
                    created = False
                row["status"] = "created" if created else "exists"
                created_count += int(created)
            results.append(row)
        return Response({
            "confirmed": confirm,
            "total": len(results),
            "valid": sum(r["status"] in {"valid", "created", "exists"} for r in results),
            "with_errors": sum(r["status"] == "error" for r in results),
            "created": created_count,
            "results": results,
        })
