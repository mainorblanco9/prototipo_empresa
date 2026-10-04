from io import BytesIO
import json
from decimal import Decimal
from unittest.mock import MagicMock, patch

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook

from .ai_client import AIProviderError, generate_local_text, generate_text
from .invoice_documents import InvoiceMappingError, build_invoice_batch
from .issued_invoice import (
    IssuedInvoiceDraftError,
    build_issued_invoice_draft,
)
from .local_simulation import simulate_local_invoice_submission
from .models import InvoiceAuditEvent, UserDesignPreference
from .sharepoint import is_sharepoint_configured, upload_to_sharepoint
from .table_analysis import analyze_table
from .validation import (
    FileValidationError,
    _parse_decimal,
    build_csv_report,
    validate_upload,
)


class DynamicTableAnalysisTests(TestCase):
    def monthly_workbook_upload(self):
        workbook = Workbook()
        empty_sheet = workbook.active
        empty_sheet.title = "Portada"
        empty_sheet["A1"] = "Libro financiero"
        sheet = workbook.create_sheet("Flujo anual")
        sheet.merge_cells("B1:F1")
        sheet["B1"] = "Flujo de caja mensual"
        sheet.append([None, "CONCEPTO", "ENERO", "FEBRERO", "MARZO", "TOTAL ANUAL"])
        sheet.append([None, "RESUMEN DE EFECTIVO"])
        sheet.append([None, "Flujo Neto del Período", 100, -25, 75, 150])
        sheet.append([None, "Saldo Final Acumulado", 100, 75, 150, 150])
        sheet.append([None, "RESUMEN AÑO 1"])
        sheet.append([None, "Flujo Neto Anual", 1000])
        next_sheet = workbook.create_sheet("AÑO 2")
        next_sheet.append([None, "CONCEPTO", "ENERO", "FEBRERO", "MARZO", "TOTAL ANUAL"])
        next_sheet.append([None, "Flujo Neto del Período", 200, 125, 75, 400])
        next_sheet.append([None, "Saldo Final Acumulado", 200, 325, 400, 400])
        content = BytesIO()
        workbook.save(content)
        workbook.close()
        return SimpleUploadedFile(
            "flujo.xlsx",
            content.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def row_oriented_workbook_upload(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Flujo de caja"
        sheet.append(
            [
                "Mes",
                "Hospedaje",
                "Comisiones Tours",
                "Comisiones Artesanías",
                "Ingresos Totales",
                "Costos Totales",
                "Flujo Neto",
            ]
        )
        months = (
            "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
            "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
        )
        for row_number, month in enumerate(months, start=2):
            sheet.append(
                [
                    month,
                    2_160_000,
                    900_000,
                    540_000,
                    f"=B{row_number}+C{row_number}+D{row_number}",
                    f"=600000+900000+250000+400000+200000+150000",
                    f"=E{row_number}-F{row_number}",
                ]
            )
        sheet.append(["TOTAL ANUAL", None, None, None, "=SUM(E2:E13)"])
        content = BytesIO()
        workbook.save(content)
        workbook.close()
        return SimpleUploadedFile(
            "flujo_profesor.xlsx",
            content.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True, AI_PROVIDER="disabled")
    @patch(
        "validador.table_analysis.generate_text",
        return_value='{"label_index":0,"period_indices":[1,2,3]}',
    )
    def test_reads_merged_title_and_finds_monthly_matrix_on_non_active_sheet(
        self, generate_text
    ):
        analysis = analyze_table(self.monthly_workbook_upload())

        self.assertEqual(analysis["source_sheet"], "Flujo anual")
        self.assertEqual(analysis["headers"][:4], ["Concepto", "Enero", "Febrero", "Marzo"])
        self.assertEqual(len(analysis["rows"]), 2)
        self.assertEqual(analysis["matrix_chart"]["selected"], "Flujo Neto del Período")
        self.assertEqual(
            [point["label"] for point in analysis["chart"]["points"]],
            ["Enero", "Febrero", "Marzo"],
        )
        self.assertEqual(
            [point["value"] for point in analysis["chart"]["points"]],
            [100.0, -25.0, 75.0],
        )
        self.assertEqual(
            [sheet["sheet_name"] for sheet in analysis["matrix_chart"]["sheets"]],
            ["Flujo anual", "AÑO 2"],
        )
        prompt = generate_text.call_args.kwargs["user_prompt"]
        self.assertNotIn("Flujo Neto del Período", prompt)
        self.assertNotIn("-25", prompt)

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True, AI_PROVIDER="disabled")
    @patch(
        "validador.table_analysis.generate_text",
        return_value='{"label_index":0,"period_indices":[1,2,3]}',
    )
    def test_monthly_chart_is_selectable_without_a_second_ai_call(self, _generate_text):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        response = self.client.post(
            reverse("index"),
            {"archivo": self.monthly_workbook_upload()},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Indicador para comparar por mes")
        self.assertContains(response, "Hoja para el gráfico")
        self.assertContains(response, "AÑO 2")
        self.assertContains(response, "Saldo Final Acumulado")
        self.assertContains(response, "sin otra consulta a la IA")
        self.assertEqual(response.context["source_sheet"], "Flujo anual")

    @patch(
        "validador.table_analysis.generate_text",
        return_value='{"x_index":0,"y_index":6,"aggregation":"sum"}',
    )
    def test_analyzes_months_in_rows_and_calculates_excel_formulas(self, _generate_text):
        analysis = analyze_table(self.row_oriented_workbook_upload())

        self.assertEqual(analysis["source_sheet"], "Flujo de caja")
        self.assertEqual(len(analysis["rows"]), 12)
        self.assertEqual(analysis["rows"][0]["Ingresos Totales"], "3600000")
        self.assertEqual(analysis["rows"][0]["Costos Totales"], "2500000")
        self.assertEqual(analysis["rows"][0]["Flujo Neto"], "1100000")
        self.assertEqual(analysis["matrix_chart"]["selected"], "Flujo Neto")
        self.assertEqual(analysis["chart"]["points"][0]["label"], "Enero")
        self.assertEqual(analysis["chart"]["points"][0]["value"], 1100000.0)
        self.assertEqual(
            [point["label"] for point in analysis["chart"]["points"]],
            [
                "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
                "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
            ],
        )

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True, AI_PROVIDER="disabled")
    @patch(
        "validador.table_analysis.generate_text",
        return_value='{"x_index":0,"y_index":6,"aggregation":"sum"}',
    )
    def test_upload_page_offers_monthly_measure_selector_for_row_oriented_excel(
        self, _generate_text
    ):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        response = self.client.post(
            reverse("index"),
            {"archivo": self.row_oriented_workbook_upload()},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Indicador para comparar por mes")
        self.assertContains(response, "Flujo Neto")
        self.assertContains(response, "Ingresos Totales")
        self.assertEqual(response.context["matrix_chart"]["selected"], "Flujo Neto")

    @patch(
        "validador.table_analysis.generate_text",
        return_value='{"x_index":0,"y_index":1,"aggregation":"sum"}',
    )
    def test_uses_ollama_schema_recommendation_to_group_numeric_data(self, generate_text):
        upload = SimpleUploadedFile(
            "ventas.csv",
            (
                "Región;Ventas\n"
                "Norte;120\n"
                "Sur;90\n"
                "Norte;30\n"
            ).encode("utf-8"),
        )

        analysis = analyze_table(upload)

        self.assertEqual(analysis["headers"], ["Region", "Ventas"])
        self.assertEqual(
            analysis["chart"]["points"],
            [
                {"label": "Norte", "value": 150.0, "width": 100.0},
                {"label": "Sur", "value": 90.0, "width": 60.0},
            ],
        )
        prompt = generate_text.call_args.kwargs["user_prompt"]
        self.assertIn('"type": "number"', prompt)
        self.assertNotIn("Norte", prompt)
        self.assertNotIn("120", prompt)

    @patch(
        "validador.table_analysis.generate_text",
        return_value='{"x_index":0,"y_index":-1,"aggregation":"count"}',
    )
    def test_creates_a_frequency_chart_for_text_only_tables(self, _generate_text):
        upload = SimpleUploadedFile(
            "estados.tsv",
            "Estado\tEstado\nActivo\tActivo\nInactivo\tInactivo\nActivo\tActivo\n".encode(
                "utf-8"
            ),
        )

        analysis = analyze_table(upload)

        self.assertEqual(analysis["headers"], ["Estado", "Estado 2"])
        self.assertEqual(
            analysis["chart"]["points"],
            [
                {"label": "Activo", "value": 2, "width": 100.0},
                {"label": "Inactivo", "value": 1, "width": 50.0},
            ],
        )

    @patch(
        "validador.table_analysis.generate_text",
        return_value='{"x_index":-1,"y_index":1,"aggregation":"sum"}',
    )
    def test_uses_detected_category_when_ai_selects_row_index(
        self, _generate_text
    ):
        upload = SimpleUploadedFile(
            "ventas.csv",
            b"Region,Sales\nNorth,120\nSouth,90\nNorth,30\n",
        )

        analysis = analyze_table(upload)

        self.assertEqual(analysis["chart"]["x_label"], "Region")
        self.assertEqual(
            analysis["chart"]["points"][0]["label"],
            "North",
        )

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True, AI_PROVIDER="disabled")
    def test_upload_page_displays_dynamic_columns_and_chart_for_general_data(self):
        client = self.client
        client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        response = client.post(
            reverse("index"),
            {
                "archivo": SimpleUploadedFile(
                    "ventas.csv",
                    b"Region,Sales\nNorth,120\nSouth,90\n",
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Suma de Sales por Region")
        self.assertContains(response, "North")
        self.assertContains(response, "Sales")
        self.assertContains(response, "Esta tabla no es una plantilla de gastos")
        self.assertFalse(response.context["expense_validation"])
        report = client.get(response.context["report_url"])
        self.assertEqual(report.status_code, 200)
        report_text = report.content.decode("utf-8-sig")
        self.assertIn("Region,Sales", report_text)
        self.assertIn("North,120", report_text)


@override_settings(
    AI_PROVIDER="ollama",
    OLLAMA_BASE_URL="http://localhost:11434",
    OLLAMA_MODEL="qwen2.5:3b",
    OLLAMA_TIMEOUT_SECONDS=90,
)
class LocalAIClientTests(TestCase):
    @patch("validador.ai_client.requests.post")
    def test_sends_prompts_to_local_ollama(self, post):
        response = MagicMock()
        response.json.return_value = {"message": {"content": "abstract"}}
        post.return_value = response

        result = generate_text("system prompt", "user prompt", max_tokens=10)

        self.assertEqual(result, "abstract")
        post.assert_called_once_with(
            "http://localhost:11434/api/chat",
            json={
                "model": "qwen2.5:3b",
                "messages": [
                    {"role": "system", "content": "system prompt"},
                    {"role": "user", "content": "user prompt"},
                ],
                "stream": False,
                "options": {"num_predict": 10},
            },
            timeout=90,
        )

    @patch("validador.ai_client.requests.post")
    def test_requests_json_mode_from_ollama(self, post):
        response = MagicMock()
        response.json.return_value = {"message": {"content": '{"x_index":0}'}}
        post.return_value = response

        generate_text("system prompt", "user prompt", json_mode=True)

        self.assertEqual(
            post.call_args.kwargs["json"]["format"],
            "json",
        )

    @patch("validador.ai_client.requests.post")
    def test_reports_when_local_ollama_is_unavailable(self, post):
        import requests

        post.side_effect = requests.ConnectionError("connection refused")

        with self.assertRaisesMessage(AIProviderError, "No se pudo conectar con Ollama"):
            generate_text("system prompt", "user prompt")

    @override_settings(AI_PROVIDER="openai", OPENAI_API_KEY="test-key")
    @patch("validador.ai_client.OpenAI")
    def test_openai_remains_available_as_an_alternative(self, openai):
        openai.return_value.responses.create.return_value.output_text = "minimal"

        result = generate_text("system prompt", "user prompt", max_tokens=10)

        self.assertEqual(result, "minimal")
        openai.assert_called_once_with(api_key="test-key")
        openai.return_value.responses.create.assert_called_once_with(
            model="gpt-4.1-mini",
            instructions="system prompt",
            input="user prompt",
            max_output_tokens=10,
        )

    @override_settings(
        AI_PROVIDER="openai",
        OLLAMA_BASE_URL="http://localhost:11434",
        OLLAMA_MODEL="qwen2.5:3b",
        OLLAMA_TIMEOUT_SECONDS=90,
    )
    @patch("validador.ai_client.requests.post")
    def test_local_generation_stays_on_ollama_even_if_openai_is_selected(self, post):
        response = MagicMock()
        response.json.return_value = {"message": {"content": '{"ok":true}'}}
        post.return_value = response

        result = generate_local_text("system", "local data", json_mode=True)

        self.assertEqual(result, '{"ok":true}')
        self.assertEqual(
            post.call_args.args[0],
            "http://localhost:11434/api/chat",
        )

    @override_settings(
        AI_PROVIDER="ollama",
        OLLAMA_BASE_URL="http://localhost:11434",
        OLLAMA_MODEL="qwen2.5:3b",
        OLLAMA_TIMEOUT_SECONDS=90,
    )
    @patch("validador.ai_client.requests.post")
    def test_local_generation_passes_json_schema_to_ollama(self, post):
        response = MagicMock()
        response.json.return_value = {
            "message": {"content": '{"mapping":{"invoice":"folio"}}'}
        }
        post.return_value = response
        schema = {
            "type": "object",
            "properties": {
                "mapping": {
                    "type": "object",
                    "properties": {
                        "invoice": {"type": "string", "enum": ["folio"]}
                    },
                }
            },
        }

        generate_local_text("system", "local data", json_schema=schema)

        self.assertEqual(post.call_args.kwargs["json"]["format"], schema)

    @override_settings(OLLAMA_BASE_URL="https://ollama.example.com")
    def test_local_generation_rejects_remote_ollama_urls(self):
        with self.assertRaisesMessage(AIProviderError, "debe apuntar a localhost"):
            generate_local_text("system", "private invoice data")


@override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True, AI_PROVIDER="disabled")
class ExpenseValidationTests(TestCase):
    def test_parses_common_localized_invoice_amount_formats(self):
        values = {
            "1.234,56": Decimal("1234.56"),
            "1,234.56": Decimal("1234.56"),
            "1.234.567,89": Decimal("1234567.89"),
            "1,234,567.89": Decimal("1234567.89"),
            "CRC 1.234,56": Decimal("1234.56"),
            "₡1.234,56": Decimal("1234.56"),
            "(1.234,56)": Decimal("-1234.56"),
            "1,25": Decimal("1.25"),
        }

        for raw_value, expected in values.items():
            with self.subTest(raw_value=raw_value):
                self.assertEqual(_parse_decimal(raw_value), expected)
    def upload(self, content, name="gastos.csv"):
        return SimpleUploadedFile(name, content.encode("utf-8"))

    def setUp(self):
        response = self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        self.assertEqual(response.status_code, 302)

    def test_validates_a_correct_expense_row(self):
        rows, issues = validate_upload(
            self.upload(
                "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                "2026-09-01,Proveedor A,F-001,CRC,1000,160,1160\n"
            )
        )

        self.assertEqual(issues, {})
        self.assertEqual(rows[0]["estado"], "Correcto")

    def test_rejects_foreign_currency_codes_for_costa_rica_invoice_flow(self):
        for currency in ("USD", "MXN", "EUR"):
            with self.subTest(currency=currency):
                rows, issues = validate_upload(
                    self.upload(
                        "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                        f"2026-09-01,Proveedor,F-{currency},{currency},100,13,113\n"
                    )
                )

                self.assertEqual(rows[0]["estado"], "Revisar")
                self.assertIn("Moneda no admitida", rows[0]["errores"])
                self.assertEqual(
                    issues[
                        "Moneda no admitida: solo se aceptan colones de Costa Rica (CRC)"
                    ],
                    1,
                )

    def test_reports_duplicate_invoice_and_invalid_total(self):
        rows, issues = validate_upload(
            self.upload(
                "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                "2026-09-01,Proveedor A,F-001,CRC,1000,160,1160\n"
                "2026-09-02,Proveedor A,F-001,CRC,1000,160,1200\n"
            )
        )

        self.assertEqual(rows[1]["estado"], "Revisar")
        self.assertIn("Factura duplicada", rows[1]["errores"])
        self.assertEqual(issues["Total no coincide con subtotal + impuesto"], 1)

    def test_rejects_missing_columns(self):
        with self.assertRaisesMessage(FileValidationError, "Faltan columnas requeridas"):
            validate_upload(self.upload("fecha,proveedor\n2026-09-01,Proveedor A\n"))

    def test_report_escapes_spreadsheet_formulas_in_text_fields(self):
        rows, _ = validate_upload(
            self.upload(
                "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                '2026-09-01,=HYPERLINK("https://example.invalid"),F-001,CRC,1000,160,1160\n'
            )
        )

        report = build_csv_report(rows)
        self.assertIn("'=HYPERLINK", report)

    def test_upload_page_shows_report_and_download(self):
        response = self.client.post(
            reverse("index"),
            {
                "archivo": self.upload(
                    "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                    "2026-09-01,Proveedor A,F-001,CRC,1000,160,1160\n"
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Resumen de validación")
        self.assertContains(response, "Correcto")
        self.assertContains(response, "Descargar CSV")
        report_url = response.context["report_url"]

        download = self.client.get(report_url)
        self.assertEqual(download.status_code, 200)
        self.assertIn("Correcto", download.content.decode("utf-8-sig"))

    def test_invoice_upload_keeps_invalid_date_rows_visible_and_validates_them(self):
        response = self.client.post(
            reverse("index"),
            {
                "archivo": self.upload(
                    "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                    "2026-09-01,Proveedor A,F-001,CRC,1000,160,1160\n"
                    "2026-09-02,Proveedor B,F-002,CRC,1000,160,1160\n"
                    "2026-09-03,Proveedor C,F-003,CRC,1000,160,1160\n"
                    "fecha-invalida,Proveedor D,F-004,CRC,1000,160,1160\n"
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total_rows"], 4)
        self.assertEqual(response.context["valid_rows"], 3)
        self.assertEqual(response.context["invalid_rows"], 1)
        self.assertEqual(len(response.context["rows"]), 4)
        self.assertContains(response, "Fecha inválida")

    def test_converts_validated_expense_rows_to_versioned_invoice_objects(self):
        upload = SimpleUploadedFile(
            "facturas.xlsx",
            self.excel_expense_file(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        batch = build_invoice_batch(upload)

        self.assertEqual(batch["schema_version"], "1.1.0")
        self.assertEqual(batch["currency_scope"], "CRC")
        self.assertEqual(batch["status"], "ready")
        self.assertEqual(batch["summary"]["validated_count"], 1)
        self.assertEqual(
            batch["invoices"][0],
            {
                "row_number": 2,
                "invoice_number": "FAC-1001",
                "supplier": "Papelería Central",
                "issue_date": "2026-09-01",
                "currency": "CRC",
                "subtotal": "1000",
                "tax": "160",
                "total": "1160",
                "status": "validated",
                "errors": [],
            },
        )
        self.assertEqual(
            batch["integration"]["status"],
            "awaiting_official_api_contract",
        )

    def test_mock_submission_marks_valid_batches_as_local_simulations(self):
        batch = build_invoice_batch(
            SimpleUploadedFile(
                "facturas.xlsx",
                self.excel_expense_file(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        )

        result = simulate_local_invoice_submission(batch)

        self.assertEqual(result["provider"], "Simulador local de flujo TRIBU-CR")
        self.assertEqual(result["mode"], "simulation_only")
        self.assertFalse(result["real_hacienda_contacted"])
        self.assertFalse(result["real_submission"])
        self.assertEqual(result["status"], "simulation_passed")
        self.assertEqual(result["simulated_ready_count"], 1)
        self.assertTrue(result["simulation_id"].startswith("LOCAL-SIM-"))
        self.assertNotIn("mock_receipt_id", result["invoices"][0])

    def test_mock_submission_blocks_batches_with_validation_errors(self):
        batch = build_invoice_batch(
            SimpleUploadedFile(
                "facturas.csv",
                (
                    "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                    "2026-09-01,Tienda,F-1,CRC,100,13,999\n"
                ).encode("utf-8"),
            )
        )

        result = simulate_local_invoice_submission(batch)

        self.assertEqual(result["status"], "simulation_blocked")
        self.assertEqual(result["simulated_ready_count"], 0)
        self.assertEqual(result["blocked_count"], 1)

    def test_mock_submission_never_accepts_a_foreign_currency_even_in_a_ready_batch(self):
        result = simulate_local_invoice_submission(
            {
                "status": "ready",
                "invoices": [
                    {
                        "invoice_number": "F-USD",
                        "currency": "USD",
                        "status": "validated",
                        "errors": [],
                    }
                ],
            }
        )

        self.assertEqual(result["status"], "simulation_blocked")
        self.assertEqual(result["simulated_ready_count"], 0)

    def test_empty_invoice_file_cannot_be_marked_ready_or_accepted_by_mock(self):
        upload = SimpleUploadedFile(
            "facturas.csv",
            (
                "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
            ).encode("utf-8"),
        )
        batch = build_invoice_batch(upload)
        result = simulate_local_invoice_submission(batch)

        self.assertEqual(batch["status"], "needs_review")
        self.assertEqual(result["status"], "simulation_blocked")
        self.assertEqual(result["simulated_ready_count"], 0)

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_pwa_assets_are_public_and_do_not_cache_invoice_pages(self):
        manifest = self.client.get(reverse("pwa_manifest"))
        self.assertEqual(manifest.status_code, 200)
        data = json.loads(manifest.content)
        self.assertEqual(data["display"], "standalone")
        self.assertEqual(data["start_url"], "/facturas/movil/")

        worker = self.client.get(reverse("pwa_service_worker"))
        self.assertEqual(worker.status_code, 200)
        self.assertEqual(worker["Service-Worker-Allowed"], "/")
        self.assertIn("no-cache", worker["Cache-Control"])
        self.assertNotIn(b"/facturas/emitidas", worker.content)
        self.assertNotIn(b"/api/", worker.content)

        self.assertEqual(self.client.get(reverse("pwa_icon")).status_code, 200)
        self.assertEqual(self.client.get(reverse("pwa_offline")).status_code, 200)
        login = Client().get(reverse("login"))
        self.assertContains(login, 'rel="manifest"')


    def test_builds_non_fiscal_issued_invoice_draft_with_crc_line_totals(self):
        draft = build_issued_invoice_draft(
            {
                "issuer_name": "Empresa Emisora",
                "issuer_id": "3-101-123456",
                "receiver_name": "Cliente Receptor",
                "receiver_id": "1-2345-0678",
                "issue_date": "2026-10-03",
            },
            [
                {
                    "description": "Asesoría",
                    "quantity": "2",
                    "unit_price": "1130",
                    "tax_rate": "13",
                },
                {
                    "description": "Hospedaje",
                    "quantity": "1",
                    "unit_price": "2825",
                    "tax_rate": "13",
                },
            ],
        )

        self.assertEqual(draft["document_type"], "factura_electronica")
        self.assertEqual(draft["status"], "internal_draft")
        self.assertEqual(draft["currency"], "CRC")
        self.assertEqual(draft["totals"], {
            "subtotal": "4500.00",
            "tax": "585.00",
            "total": "5085.00",
        })
        self.assertEqual(draft["schema_version"], "0.2.0")
        self.assertEqual(draft["lines"][0]["unit_price_including_tax"], "1130.00")
        self.assertEqual(draft["lines"][0]["subtotal"], "2000.00")
        self.assertEqual(draft["lines"][0]["tax_amount"], "260.00")
        self.assertTrue(draft["not_official_format"])
        self.assertFalse(draft["transmission"]["real_hacienda_contacted"])
        self.assertEqual(
            draft["validation_scope"],
            "internal_required_fields_arithmetic_and_inclusive_tax_calculation",
        )

    def test_rejects_invalid_or_incomplete_issued_invoice_drafts(self):
        with self.assertRaisesMessage(IssuedInvoiceDraftError, "identificación"):
            build_issued_invoice_draft(
                {
                    "issuer_name": "Empresa",
                    "issuer_id": "",
                    "receiver_name": "Cliente",
                    "receiver_id": "123",
                    "issue_date": "2026-10-03",
                },
                [
                    {
                        "description": "Servicio",
                        "quantity": "1",
                        "unit_price": "100",
                        "tax_rate": "13",
                    }
                ],
            )

        with self.assertRaisesMessage(IssuedInvoiceDraftError, "mayor que cero"):
            build_issued_invoice_draft(
                {
                    "issuer_name": "Empresa",
                    "issuer_id": "ID",
                    "receiver_name": "Cliente",
                    "receiver_id": "ID-2",
                    "issue_date": "2026-10-03",
                },
                [
                    {
                        "description": "Servicio",
                        "quantity": "0",
                        "unit_price": "100",
                        "tax_rate": "13",
                    }
                ],
            )

        with self.assertRaisesMessage(IssuedInvoiceDraftError, "tarifa de IVA"):
            build_issued_invoice_draft(
                {
                    "issuer_name": "Empresa",
                    "issuer_id": "ID",
                    "receiver_name": "Cliente",
                    "receiver_id": "ID-2",
                    "issue_date": "2026-10-03",
                },
                [
                    {
                        "description": "Servicio",
                        "quantity": "1",
                        "unit_price": "1130",
                        "tax_rate": "99",
                    }
                ],
            )

    def test_included_tax_is_extracted_for_reduced_zero_and_exempt_rates(self):
        base_data = {
            "issuer_name": "Empresa",
            "issuer_id": "ID",
            "receiver_name": "Cliente",
            "receiver_id": "ID-2",
            "issue_date": "2026-10-03",
        }

        for rate, expected_tax in (
            ("4", "38.46"),
            ("2", "19.61"),
            ("1", "9.90"),
            ("0", "0.00"),
            ("exempt", "0.00"),
        ):
            with self.subTest(rate=rate):
                draft = build_issued_invoice_draft(
                    base_data,
                    [
                        {
                            "description": "Servicio",
                            "quantity": "1",
                            "unit_price": "1000",
                            "tax_rate": rate,
                        }
                    ],
                )

                self.assertEqual(draft["lines"][0]["tax_amount"], expected_tax)
                self.assertEqual(
                    Decimal(draft["lines"][0]["subtotal"])
                    + Decimal(draft["lines"][0]["tax_amount"]),
                    Decimal("1000.00"),
                )

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_issued_invoice_draft_is_a_separate_local_flow_and_is_not_logged_as_invoice_data(self):
        response = self.client.post(
            reverse("issued_invoice_draft"),
            {
                "action": "prepare",
                "issuer_name": "Empresa Emisora Privada",
                "issuer_id": "ID-EMISOR-PRIVADO",
                "receiver_name": "Cliente Receptor Privado",
                "receiver_id": "ID-RECEPTOR-PRIVADO",
                "issue_date": "2026-10-03",
                "line_description": ["Servicio local"],
                "line_quantity": ["1"],
                "line_unit_price": ["1130"],
                "line_tax_rate": ["13"],
            },
        )

        self.assertEqual(response.status_code, 200)
        draft = response.context["draft"]
        self.assertEqual(draft["totals"]["total"], "1130.00")
        self.assertEqual(draft["totals"]["subtotal"], "1000.00")
        self.assertEqual(draft["totals"]["tax"], "130.00")
        self.assertContains(response, "no es una factura electrónica válida ni se remite a Hacienda")
        self.assertContains(response, "Precio unitario con IVA")
        self.assertContains(response, "IVA incluido (calculado)")
        self.assertContains(response, 'name="line_tax_rate"')
        self.assertContains(response, "Descargar JSON interno")
        self.assertIsNone(response.context.get("batch"))
        event = InvoiceAuditEvent.objects.get(
            batch_id=draft["draft_id"],
            event_type=InvoiceAuditEvent.EventType.ISSUED_DRAFT,
        )
        self.assertEqual(event.result, "draft_prepared")
        self.assertFalse(hasattr(event, "issuer_name"))
        self.assertFalse(hasattr(event, "receiver_id"))

        downloaded = self.client.get(
            reverse("download_issued_invoice_draft")
        )
        self.assertEqual(downloaded.status_code, 200)
        self.assertFalse(
            json.loads(downloaded.content)["transmission"]["real_hacienda_contacted"]
        )
        audit = self.client.get(reverse("invoice_audit"))
        self.assertNotContains(audit, "Empresa Emisora Privada")
        self.assertNotContains(audit, "Servicio local")
        self.assertContains(audit, "borrador interno · no fiscal")

    @patch("validador.invoice_documents.generate_local_text")
    def test_monthly_cash_flow_workbook_is_rejected_as_not_an_invoice_file(
        self, generate_local_text
    ):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "AÑO 1"
        sheet.append(["Flujo de caja"])
        sheet.append(
            [
                "CONCEPTO",
                "ENERO",
                "FEBRERO",
                "MARZO",
                "ABRIL",
                "MAYO",
                "JUNIO",
                "JULIO",
                "AGOSTO",
                "SEPTIEMBRE",
                "OCTUBRE",
                "NOVIEMBRE",
                "DICIEMBRE",
            ]
        )
        sheet.append(["Ingresos", 100, 100, 100, 100, 100, 100, 100, 100, 100, 100, 100, 100])
        content = BytesIO()
        workbook.save(content)
        workbook.close()
        upload = SimpleUploadedFile(
            "flujo_de_caja.xlsx",
            content.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

        with self.assertRaisesMessage(
            InvoiceMappingError,
            "no un archivo de facturas",
        ):
            build_invoice_batch(upload)

        generate_local_text.assert_not_called()

    @patch(
        "validador.invoice_documents.generate_local_text",
        return_value=(
            '{"table_index":0,"mapping":{"fecha":"fecha_emision",'
            '"proveedor":"nombre_comercio","numero_factura":"folio_fiscal",'
            '"moneda":"divisa","subtotal":"base_imponible",'
            '"impuesto":"iva_cobrado","total":"importe_pagado"}}'
        ),
    )
    def test_uses_local_ai_to_map_unfamiliar_headers_with_local_sample_values(
        self, generate_local_text
    ):
        upload = SimpleUploadedFile(
            "facturas.csv",
            (
                "Fecha emisión,Nombre comercio,Folio fiscal,Divisa,"
                "Base imponible,IVA cobrado,Importe pagado\n"
                "2026-09-01,Tienda Secreta,F-77,CRC,10000,1300,11300\n"
            ).encode("utf-8"),
        )

        batch = build_invoice_batch(upload)

        self.assertEqual(batch["status"], "ready")
        self.assertEqual(batch["mapping_source"], "Ollama local")
        self.assertEqual(batch["invoices"][0]["invoice_number"], "F-77")
        self.assertEqual(batch["invoices"][0]["total"], "11300")
        prompt = generate_local_text.call_args.kwargs["user_prompt"]
        self.assertIn("fecha_emision", prompt)
        self.assertIn("Tienda Secreta", prompt)
        self.assertIn("11300", prompt)
        self.assertIn(
            "este Ollama local",
            generate_local_text.call_args.kwargs["system_prompt"],
        )
        self.assertIn(
            "nombre de encabezado exacto",
            generate_local_text.call_args.kwargs["system_prompt"],
        )

    @patch(
        "validador.invoice_documents.generate_local_text",
        side_effect=[
            '{"table_index":0,"mapping":{"fecha":12}}',
            (
                '{"table_index":0,"mapping":{"fecha":0,"proveedor":1,'
                '"numero_factura":2,"moneda":3,"subtotal":4,'
                '"impuesto":5,"total":6}}'
            ),
        ],
    )
    def test_retries_mapping_when_ollama_returns_column_index_out_of_range(
        self, generate_local_text
    ):
        upload = SimpleUploadedFile(
            "facturas.csv",
            b"Dia,Comercio,Folio,Divisa,Base,IVA,Monto\n"
            b"2026-09-01,Tienda,F-101,CRC,10000,1300,11300\n",
        )

        batch = build_invoice_batch(upload)

        self.assertEqual(batch["status"], "ready")
        self.assertEqual(batch["invoices"][0]["invoice_number"], "F-101")
        self.assertEqual(generate_local_text.call_count, 2)
        retry_prompt = json.loads(
            generate_local_text.call_args_list[1].kwargs["user_prompt"]
        )
        self.assertIn("columna inexistente", retry_prompt["correction"])
        self.assertIn('"dia"', retry_prompt["correction"])

    @patch(
        "validador.invoice_documents.generate_local_text",
        side_effect=[
            '{"table_index":0,"mapping":{"fecha":"fecha_que_no_existe"}}',
            (
                '{"table_index":0,"mapping":{"fecha":"dia",'
                '"proveedor":"comercio","numero_factura":"folio",'
                '"moneda":"divisa","subtotal":"base","impuesto":"iva",'
                '"total":"monto"}}'
            ),
        ],
    )
    def test_retries_when_ollama_returns_column_name_not_in_table(
        self, generate_local_text
    ):
        upload = SimpleUploadedFile(
            "facturas.csv",
            b"Dia,Comercio,Folio,Divisa,Base,IVA,Monto\n"
            b"2026-09-01,Tienda,F-102,CRC,10000,1300,11300\n",
        )

        batch = build_invoice_batch(upload)

        self.assertEqual(batch["status"], "ready")
        self.assertEqual(batch["invoices"][0]["invoice_number"], "F-102")
        self.assertEqual(generate_local_text.call_count, 2)
        retry_prompt = json.loads(
            generate_local_text.call_args_list[1].kwargs["user_prompt"]
        )
        self.assertIn('"dia"', retry_prompt["correction"])
        self.assertIn("fecha_que_no_existe", retry_prompt["correction"])

    @patch(
        "validador.invoice_documents.generate_local_text",
        return_value='{"table_index":0,"mapping":{"fecha":0}}',
    )
    def test_rejects_ai_header_mapping_when_required_fields_are_missing(
        self, _generate_local_text
    ):
        upload = SimpleUploadedFile(
            "facturas.csv",
            b"Dia,Comercio,Folio,Divisa,Base,IVA,Monto\n"
            b"2026-09-01,Tienda,F-77,CRC,10000,1300,11300\n",
        )

        with self.assertRaisesMessage(InvoiceMappingError, "No se reconocieron"):
            build_invoice_batch(upload)

    @patch(
        "validador.invoice_documents.generate_local_text",
        return_value=(
            '{"table_index":0,"mapping":{"fecha":100,"proveedor":101,'
            '"numero_factura":102,"moneda":103,"subtotal":104,'
            '"impuesto":105,"total":106}}'
        ),
    )
    def test_maps_invoice_columns_beyond_the_old_column_limit(
        self, generate_local_text
    ):
        headers = [f"campo_{index}" for index in range(107)]
        values = [""] * len(headers)
        values[100:107] = [
            "2026-09-01",
            "Tienda",
            "F-100",
            "CRC",
            "10000",
            "1300",
            "11300",
        ]
        content = ",".join(headers) + "\n" + ",".join(values) + "\n"

        batch = build_invoice_batch(
            SimpleUploadedFile("facturas.csv", content.encode("utf-8"))
        )

        self.assertEqual(batch["status"], "ready")
        self.assertEqual(batch["invoices"][0]["invoice_number"], "F-100")
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(headers)
        worksheet.append(values)
        excel_content = BytesIO()
        workbook.save(excel_content)
        workbook.close()
        excel_batch = build_invoice_batch(
            SimpleUploadedFile(
                "facturas.xlsx",
                excel_content.getvalue(),
                content_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
            )
        )

        self.assertEqual(excel_batch["status"], "ready")
        self.assertEqual(excel_batch["invoices"][0]["invoice_number"], "F-100")
        for call in generate_local_text.call_args_list:
            prompt = json.loads(call.kwargs["user_prompt"])
            self.assertEqual(len(prompt["tables"][0]["columns"]), 107)
            self.assertEqual(prompt["tables"][0]["columns"][-1]["index"], 106)

    @patch(
        "validador.invoice_documents.generate_local_text",
        return_value=(
            '{"table_index":0,"mapping":{"fecha":0,"proveedor":1,'
            '"numero_factura":2,"moneda":3,"subtotal":4,'
            '"impuesto":5,"total":6}}'
        ),
    )
    def test_uses_variable_sample_row_count_based_on_table_width(
        self, generate_local_text
    ):
        rows = [
            f"2026-09-{day:02d},Tienda,F-{day},CRC,100,13,113"
            for day in range(1, 12)
        ]
        upload = SimpleUploadedFile(
            "facturas.csv",
            (
                "Dia,Comercio,Folio,Divisa,Base,IVA,Monto\n"
                + "\n".join(rows)
            ).encode("utf-8"),
        )

        batch = build_invoice_batch(upload)

        self.assertEqual(batch["summary"]["invoice_count"], 11)
        prompt = json.loads(generate_local_text.call_args.kwargs["user_prompt"])
        self.assertEqual(len(prompt["tables"][0]["sample_rows"]), 10)

    def test_processes_more_than_twenty_thousand_invoice_rows_locally(self):
        rows = (
            f"2026-09-01,Tienda,F-{number},CRC,100,13,113"
            for number in range(20_001)
        )
        content = (
            "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
            + "\n".join(rows)
        )

        batch = build_invoice_batch(
            SimpleUploadedFile("facturas.csv", content.encode("utf-8"))
        )

        self.assertEqual(batch["summary"]["invoice_count"], 20_001)
        self.assertEqual(batch["invoices"][-1]["row_number"], 20_002)

    def test_detects_csv_header_after_preamble_and_reports_source_row(self):
        upload = SimpleUploadedFile(
            "facturas.csv",
            (
                "Reporte de compras,,,,,,\n"
                "Generado localmente,,,,,,\n"
                "Fecha,Proveedor,Factura,Moneda,Subtotal,Impuesto,Total\n"
                "2026-09-01,Tienda,F-88,CRC,10000,1300,11300\n"
            ).encode("utf-8"),
        )

        batch = build_invoice_batch(upload)

        self.assertEqual(batch["status"], "ready")
        self.assertEqual(batch["header_row"], 3)
        self.assertEqual(batch["invoices"][0]["row_number"], 4)

    def test_detects_invoice_table_on_non_active_excel_sheet(self):
        workbook = Workbook()
        cover = workbook.active
        cover.title = "Portada"
        cover["A1"] = "Reporte de gastos"
        worksheet = workbook.create_sheet("Facturas")
        worksheet.append(["Exportado por sistema contable"])
        worksheet.append(
            [
                "Fecha",
                "Proveedor",
                "Factura",
                "Moneda",
                "Subtotal",
                "Impuesto",
                "Total",
            ]
        )
        worksheet.append(
            ["2026-09-01", "Tienda", "F-89", "CRC", 10000, 1300, 11300]
        )
        content = BytesIO()
        workbook.save(content)
        workbook.close()

        batch = build_invoice_batch(
            SimpleUploadedFile(
                "facturas.xlsx",
                content.getvalue(),
                content_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
            )
        )

        self.assertEqual(batch["status"], "ready")
        self.assertEqual(batch["source_sheet"], "Facturas")
        self.assertEqual(batch["header_row"], 2)
        self.assertEqual(batch["invoices"][0]["row_number"], 3)

    @override_settings(AI_PROVIDER="openai")
    def test_does_not_call_openai_when_mapping_unrecognized_headers(
        self
    ):
        upload = SimpleUploadedFile(
            "facturas.csv",
            b"Dia,Comercio,Folio,Divisa,Base,IVA,Monto\n"
            b"2026-09-01,Tienda,F-90,CRC,10000,1300,11300\n",
        )

        with (
            patch(
                "validador.ai_client._generate_with_ollama",
                return_value=(
                    '{"table_index":0,"mapping":{"fecha":0,"proveedor":1,'
                    '"numero_factura":2,"moneda":3,"subtotal":4,'
                    '"impuesto":5,"total":6}}'
                ),
            ) as generate_ollama,
            patch("validador.ai_client._generate_with_openai") as generate_openai,
        ):
            batch = build_invoice_batch(upload)

        self.assertEqual(batch["status"], "ready")
        generate_ollama.assert_called_once()
        generate_openai.assert_not_called()

    def excel_expense_file(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(
            [
                "fecha",
                "proveedor",
                "numero_factura",
                "moneda",
                "subtotal",
                "impuesto",
                "total",
            ]
        )
        worksheet.append(
            [
                "2026-09-01",
                "Papelería Central",
                "FAC-1001",
                "CRC",
                1000,
                160,
                1160,
            ]
        )
        content = BytesIO()
        workbook.save(content)
        workbook.close()
        return content.getvalue()

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_invoice_parse_api_requires_login(self):
        response = Client().post(reverse("parse_invoice_documents"))

        self.assertEqual(response.status_code, 302)
        self.assertIn("/login/", response.url)

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_invoice_parse_api_returns_structured_validated_document(self):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        response = self.client.post(
            reverse("parse_invoice_documents"),
            {
                "archivo": SimpleUploadedFile(
                    "facturas.xlsx",
                    self.excel_expense_file(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["status"], "ready")
        self.assertEqual(payload["invoices"][0]["invoice_number"], "FAC-1001")

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_local_simulation_api_does_not_create_a_fiscal_receipt(self):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        response = self.client.post(
            reverse("local_invoice_simulation"),
            {
                "archivo": SimpleUploadedFile(
                    "facturas.xlsx",
                    self.excel_expense_file(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        simulation = response.json()["simulation"]
        self.assertEqual(simulation["status"], "simulation_passed")
        self.assertFalse(simulation["real_hacienda_contacted"])
        self.assertFalse(simulation["real_submission"])
        self.assertNotIn("mock_receipt_id", simulation["invoices"][0])

    @override_settings(DEBUG=False)
    def test_local_simulation_api_is_disabled_outside_debug(self):
        user = get_user_model().objects.create_user(username="production-test")
        self.client.force_login(user)

        response = self.client.post(
            reverse("local_invoice_simulation"),
            {"archivo": SimpleUploadedFile("facturas.csv", b"")},
        )

        self.assertEqual(response.status_code, 404)
        self.assertIn("simulación técnica local", response.json()["error"])

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_sample_invoice_csv_download_contains_three_october_crc_invoices(self):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )

        response = self.client.get(reverse("download_sample_invoices"))

        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment;", response["Content-Disposition"])
        self.assertIn(b"2026-10-01", response.content)
        self.assertIn(b",CRC,", response.content)
        self.assertEqual(len(response.content.decode("utf-8").splitlines()), 4)
        parsed = self.client.post(
            reverse("invoice_workspace"),
            {
                "action": "parse",
                "archivo": SimpleUploadedFile(
                    "ejemplo_facturas_octubre.csv",
                    response.content,
                    content_type="text/csv",
                ),
            },
        )
        self.assertEqual(parsed.status_code, 200)
        self.assertEqual(parsed.context["batch"]["status"], "ready")
        self.assertEqual(parsed.context["batch"]["summary"]["invoice_count"], 3)
        self.assertEqual(parsed.context["selected_invoice_month"], "2026-10")
        chart_invoices = parsed.context["invoice_chart_data"]
        self.assertEqual(len(chart_invoices), 3)
        self.assertTrue(all(invoice["currency"] == "CRC" for invoice in chart_invoices))

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_invoice_workspace_shows_monthly_chart_and_gates_invalid_batches(self):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        parsed = self.client.post(
            reverse("invoice_workspace"),
            {
                "action": "parse",
                "archivo": SimpleUploadedFile(
                    "octubre.csv",
                    (
                        "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                        "2026-10-01,Hotel Bahía Verde,F-CR-1001,CRC,150000,19500,169500\n"
                        "2026-10-02,Restaurante Osa Norte,F-CR-1002,CRC,85000,11050,96050\n"
                        "2026-10-03,Servicios Ticos,F-CR-1003,CRC,42500,5525,48025\n"
                    ).encode("utf-8"),
                ),
            },
        )

        self.assertEqual(parsed.status_code, 200)
        self.assertContains(parsed, "Modelo de facturas interpretado")
        self.assertContains(parsed, "Análisis y validación de facturas")
        self.assertContains(parsed, "Importe por factura y mes")
        self.assertContains(parsed, "preparar modelo local para TRIBU-CR")
        self.assertContains(parsed, "2026-10")
        self.assertContains(parsed, "Campo del modelo")
        self.assertContains(parsed, "Columna detectada")
        self.assertContains(parsed, "Desliza horizontalmente")
        self.assertContains(parsed, "invoice-table")
        self.assertContains(parsed, "Descargar facturas ficticias de octubre")
        self.assertNotContains(parsed, "Simular remisión")
        self.assertNotContains(parsed, "archivo-reintentar")

        prepared = self.client.get(reverse("invoice_model_workspace"))
        self.assertEqual(prepared.status_code, 200)
        simulated = self.client.post(
            reverse("invoice_model_workspace"),
            {"action": "simulate"},
        )
        self.assertEqual(simulated.status_code, 200)
        self.assertEqual(
            simulated.context["mock_result"]["status"],
            "simulation_passed",
        )
        self.assertContains(simulated, "no se transmitió nada a Hacienda")
        self.assertContains(simulated, "Simulación local superada")
        self.assertContains(simulated, "Envío real a Hacienda")
        self.assertEqual(
            simulated.context["trace_steps"][4]["state"],
            "complete",
        )
        self.assertEqual(
            simulated.context["trace_steps"][5]["state"],
            "pending",
        )
        reloaded_model = self.client.get(reverse("invoice_model_workspace"))
        self.assertEqual(reloaded_model.context["trace_steps"][4]["state"], "complete")
        self.assertContains(reloaded_model, "Último resultado: Superada")

        invalid = self.client.post(
            reverse("invoice_workspace"),
            {
                "action": "parse",
                "archivo": SimpleUploadedFile(
                    "revision.csv",
                    (
                        "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                        "2026-10-03,Proveedor,F-4,CRC,100,13,999\n"
                    ).encode("utf-8"),
                ),
            },
        )
        self.assertEqual(invalid.status_code, 200)
        self.assertContains(invalid, "Corrígelas antes de continuar")
        self.assertNotContains(
            invalid,
            "Continuar: preparar modelo local para TRIBU-CR",
        )
        self.assertRedirects(
            self.client.get(reverse("invoice_model_workspace")),
            reverse("invoice_workspace"),
        )

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_foreign_currency_invoice_stays_out_of_crc_model_and_charts(self):
        response = self.client.post(
            reverse("invoice_workspace"),
            {
                "action": "parse",
                "archivo": SimpleUploadedFile(
                    "factura-usd.csv",
                    (
                        "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                        "2026-10-01,Proveedor,F-USD,USD,100,13,113\n"
                    ).encode("utf-8"),
                ),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["batch"]["status"], "needs_review")
        self.assertEqual(
            response.context["batch"]["invoices"][0]["currency"],
            "USD",
        )
        self.assertIn(
            "Moneda no admitida",
            response.context["batch"]["invoices"][0]["errors"][0],
        )
        self.assertFalse(response.context["model_ready"])
        self.assertEqual(response.context["invoice_currency_options"], [])
        self.assertRedirects(
            self.client.get(reverse("invoice_model_workspace")),
            reverse("invoice_workspace"),
        )

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_invoice_model_page_recovers_batch_after_local_cache_restart(self):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        self.client.post(
            reverse("invoice_workspace"),
            {
                "action": "parse",
                "archivo": SimpleUploadedFile(
                    "facturas.xlsx",
                    self.excel_expense_file(),
                    content_type=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                ),
            },
        )
        batch_id = self.client.session["invoice_workspace_batch_id"]
        cache.delete(f"invoice-workspace-batch:{batch_id}")

        response = self.client.get(reverse("invoice_model_workspace"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["batch"]["summary"]["validated_count"], 1)
        self.assertIsNotNone(cache.get(f"invoice-workspace-batch:{batch_id}"))
        analysis_page = self.client.get(reverse("invoice_workspace"))
        self.assertEqual(analysis_page.context["batch"]["status"], "ready")
        self.assertContains(analysis_page, "Preparar modelo")

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_old_unscoped_batch_is_rejected_after_crc_policy_change(self):
        legacy_batch = build_invoice_batch(
            SimpleUploadedFile(
                "facturas.csv",
                (
                    "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                    "2026-09-01,Proveedor,F-1,CRC,100,13,113\n"
                ).encode("utf-8"),
            )
        )
        legacy_batch["schema_version"] = "1.0.0"
        legacy_batch.pop("currency_scope")
        legacy_batch["invoices"][0]["currency"] = "MXN"

        session = self.client.session
        session["invoice_workspace_batch_id"] = "legacy-batch"
        session["invoice_workspace_batch"] = {
            "batch": legacy_batch,
            "expires_at": 9_999_999_999,
        }
        session.save()
        cache.set("invoice-workspace-batch:legacy-batch", legacy_batch, timeout=600)

        response = self.client.get(reverse("invoice_model_workspace"))

        self.assertRedirects(
            response,
            reverse("invoice_workspace"),
            fetch_redirect_response=False,
        )
        analysis_page = self.client.get(reverse("invoice_workspace"))
        self.assertContains(analysis_page, "reglas anteriores a la validación CRC")
        self.assertIsNone(analysis_page.context["batch"])
        self.assertIsNone(cache.get("invoice-workspace-batch:legacy-batch"))

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_invoice_model_page_shows_traceability_charts_and_downloads_same_batch(self):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        parsed = self.client.post(
            reverse("invoice_workspace"),
            {
                "action": "parse",
                "archivo": SimpleUploadedFile(
                    "facturas.xlsx",
                    self.excel_expense_file(),
                    content_type=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                ),
            },
        )
        self.assertEqual(parsed.status_code, 200)

        response = self.client.get(reverse("invoice_model_workspace"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["status_chart"][0]["value"], 1)
        self.assertEqual(response.context["currency_chart"][0]["currency"], "CRC")
        self.assertEqual(response.context["currency_chart"][0]["amount"], "₡1 160,00")
        self.assertContains(response, "Trazabilidad del proceso")
        self.assertContains(response, "Eventos registrados de este lote")
        self.assertContains(response, "solo se aceptan colones costarricenses")
        self.assertContains(response, "Envío real a Hacienda")
        self.assertContains(response, "Factura electrónica")
        self.assertContains(response, "Tiquete electrónico")
        self.assertContains(response, "Nota de crédito electrónica")
        self.assertContains(response, "resumen interno para revisión")
        self.assertContains(response, 'style="width: 100.0%"')

        downloaded = self.client.get(reverse("download_invoice_model"))
        self.assertEqual(downloaded.status_code, 200)
        self.assertIn("attachment;", downloaded["Content-Disposition"])
        self.assertEqual(
            json.loads(downloaded.content)["invoices"][0]["invoice_number"],
            "FAC-1001",
        )
        self.assertEqual(
            InvoiceAuditEvent.objects.filter(
                user_id=self.client.session.get("_auth_user_id")
            ).count(),
            3,
        )

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_invoice_audit_is_private_metadata_only_and_can_be_cleared(self):
        parsed = self.client.post(
            reverse("invoice_workspace"),
            {
                "action": "parse",
                "archivo": SimpleUploadedFile(
                    "archivo-privado.xlsx",
                    self.excel_expense_file(),
                    content_type=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                ),
            },
        )
        batch_id = parsed.context["batch"]["batch_id"]
        event = InvoiceAuditEvent.objects.get(batch_id=batch_id)

        self.assertEqual(event.event_type, InvoiceAuditEvent.EventType.ANALYSIS)
        self.assertEqual(event.result, "ready")
        self.assertEqual(event.invoice_count, 1)
        self.assertFalse(hasattr(event, "source_file"))
        self.assertFalse(hasattr(event, "invoice_number"))

        audit_response = self.client.get(reverse("invoice_audit"))
        self.assertEqual(audit_response.status_code, 200)
        self.assertContains(audit_response, "Eventos registrados")
        self.assertNotContains(audit_response, "archivo-privado.xlsx")
        self.assertNotContains(audit_response, "Papelería Central")
        self.assertNotContains(audit_response, "FAC-1001")
        self.assertNotContains(audit_response, "1,160.00")

        other_client = Client()
        other_user = get_user_model().objects.create_user(
            username="otro-usuario",
            password="otra-clave-segura",
        )
        self.assertTrue(
            other_client.login(
                username=other_user.username,
                password="otra-clave-segura",
            )
        )
        other_history = other_client.get(reverse("invoice_audit"))
        self.assertEqual(other_history.context["audit_page"].paginator.count, 0)

        cleared = self.client.post(
            reverse("invoice_audit"),
            {"action": "clear"},
            follow=True,
        )
        self.assertContains(cleared, "Bitácora borrada")
        self.assertFalse(InvoiceAuditEvent.objects.filter(batch_id=batch_id).exists())

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    @patch(
        "validador.views.generate_local_text",
        return_value="El lote está validado y pendiente de integración oficial.",
    )
    def test_invoice_model_ollama_insight_receives_only_aggregated_data(
        self, generate_local_text
    ):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        self.client.post(
            reverse("invoice_workspace"),
            {
                "action": "parse",
                "archivo": SimpleUploadedFile(
                    "facturas.xlsx",
                    self.excel_expense_file(),
                    content_type=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                ),
            },
        )

        response = self.client.post(
            reverse("invoice_model_workspace"),
            {"action": "analyze"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "El lote está validado")
        self.assertEqual(generate_local_text.call_count, 1)
        aggregate_data = json.loads(
            generate_local_text.call_args.kwargs["user_prompt"]
        )
        self.assertEqual(aggregate_data["validated_count"], 1)
        self.assertEqual(
            aggregate_data["validated_total_crc"],
            "1160",
        )
        self.assertNotIn("Papelería Central", str(aggregate_data))
        self.assertNotIn("FAC-1001", str(aggregate_data))

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_invoice_model_page_redirects_when_no_batch_is_cached(self):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )

        response = self.client.get(reverse("invoice_model_workspace"))

        self.assertRedirects(response, reverse("invoice_workspace"))

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_invoice_model_page_requires_a_validated_batch(self):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        self.client.post(
            reverse("invoice_workspace"),
            {
                "action": "parse",
                "archivo": SimpleUploadedFile(
                    "revision.csv",
                    (
                        "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                        "2026-10-03,Proveedor,F-4,CRC,100,13,999\n"
                    ).encode("utf-8"),
                ),
            },
        )

        response = self.client.get(reverse("invoice_model_workspace"))

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response["Location"], reverse("invoice_workspace"))
        landing_page = self.client.get(reverse("invoice_workspace"))
        self.assertContains(
            landing_page,
            "requieren revisión",
            status_code=200,
        )

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    @patch(
        "validador.views.generate_local_text",
        side_effect=AIProviderError("Ollama no está disponible."),
    )
    def test_invoice_model_keeps_charts_visible_when_ollama_is_unavailable(
        self, _generate_local_text
    ):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        self.client.post(
            reverse("invoice_workspace"),
            {
                "action": "parse",
                "archivo": SimpleUploadedFile(
                    "facturas.xlsx",
                    self.excel_expense_file(),
                    content_type=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                ),
            },
        )

        response = self.client.post(
            reverse("invoice_model_workspace"),
            {"action": "analyze"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["status_chart"][0]["value"], 1)
        self.assertContains(response, "Ollama no está disponible.")

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    @patch(
        "validador.invoice_documents.generate_local_text",
        side_effect=AIProviderError("No se pudo conectar con Ollama."),
    )
    def test_invoice_workspace_reports_when_local_ai_is_unavailable(
        self, _generate_text
    ):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        response = self.client.post(
            reverse("invoice_workspace"),
            {
                "archivo": SimpleUploadedFile(
                    "facturas.csv",
                    (
                        "Dia,Comercio,Folio,Divisa,Base,IVA,Monto\n"
                        "2026-09-01,Tienda Privada,F-77,CRC,10000,1300,11300\n"
                    ).encode("utf-8"),
                )
            },
        )

        self.assertEqual(response.status_code, 503)
        self.assertContains(
            response,
            "Inicia Ollama y comprueba el modelo configurado",
            status_code=503,
        )

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    @patch(
        "validador.invoice_documents.generate_local_text",
        return_value=(
            '{"table_index":0,"mapping":{"fecha":"fecha_emision",'
            '"proveedor":"nombre_comercio","numero_factura":"folio_fiscal",'
            '"moneda":"divisa","subtotal":"base_imponible",'
            '"impuesto":"iva_cobrado","total":"importe_pagado"}}'
        ),
    )
    def test_invoice_workspace_shows_ai_mapping_and_structured_output(
        self, _generate_text
    ):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        response = self.client.post(
            reverse("invoice_workspace"),
            {
                "archivo": SimpleUploadedFile(
                    "facturas.csv",
                    (
                        "Fecha emisión,Nombre comercio,Folio fiscal,Divisa,"
                        "Base imponible,IVA cobrado,Importe pagado\n"
                        "2026-09-01,Tienda,F-77,CRC,10000,1300,11300\n"
                    ).encode("utf-8"),
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ollama local")
        self.assertContains(response, "F-77")
        self.assertContains(response, "awaiting_official_api_contract")

    @override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
    def test_invoice_parse_api_returns_review_records_with_unprocessable_status(self):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        response = self.client.post(
            reverse("parse_invoice_documents"),
            {
                "archivo": SimpleUploadedFile(
                    "gastos.csv",
                    (
                        "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                        "fecha-invalida,Proveedor,F-1,CRC,100,13,113\n"
                    ).encode("utf-8"),
                )
            },
        )

        self.assertEqual(response.status_code, 422)
        payload = response.json()
        self.assertEqual(payload["status"], "needs_review")
        self.assertIn("Fecha inválida", payload["invoices"][0]["errors"])

    @override_settings(
        SHAREPOINT_TENANT_ID="tenant",
        SHAREPOINT_CLIENT_ID="client",
        SHAREPOINT_CLIENT_SECRET="secret",
        SHAREPOINT_DRIVE_ID="drive",
    )
    @patch("validador.views.upload_to_sharepoint")
    def test_uploads_original_to_sharepoint_when_all_rows_are_valid(
        self, upload_to_sharepoint
    ):
        upload_to_sharepoint.return_value = {
            "name": "gastos.csv",
            "web_url": "https://sharepoint.example/file",
            "item_id": "item-id",
        }
        uploaded_file = self.upload(
            "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
            "2026-09-01,Proveedor A,F-001,CRC,1000,160,1160\n"
        )

        response = self.client.post(reverse("index"), {"archivo": uploaded_file})

        self.assertEqual(response.status_code, 200)
        upload_to_sharepoint.assert_called_once()
        self.assertEqual(upload_to_sharepoint.call_args.args[0].name, "gastos.csv")
        self.assertEqual(upload_to_sharepoint.call_args.args[1], "gastos.csv")
        self.assertContains(response, "Archivo guardado en SharePoint")

    @override_settings(
        SHAREPOINT_TENANT_ID="tenant",
        SHAREPOINT_CLIENT_ID="client",
        SHAREPOINT_CLIENT_SECRET="secret",
        SHAREPOINT_DRIVE_ID="drive",
    )
    @patch("validador.views.upload_to_sharepoint")
    def test_does_not_upload_file_with_validation_errors(self, upload_to_sharepoint):
        response = self.client.post(
            reverse("index"),
            {
                "archivo": self.upload(
                    "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                    "fecha-invalida,Proveedor A,F-001,CRC,1000,160,1160\n"
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        upload_to_sharepoint.assert_not_called()
        self.assertContains(response, "El archivo no se subió")

    @patch("validador.views.upload_to_sharepoint")
    def test_does_not_upload_empty_file(self, upload_to_sharepoint):
        response = self.client.post(
            reverse("index"),
            {
                "archivo": self.upload(
                    "fecha,proveedor,numero_factura,moneda,subtotal,impuesto,total\n"
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        upload_to_sharepoint.assert_not_called()
        self.assertContains(response, "no contiene filas de datos")


@override_settings(
    SHAREPOINT_TENANT_ID="tenant",
    SHAREPOINT_CLIENT_ID="client",
    SHAREPOINT_CLIENT_SECRET="secret",
    SHAREPOINT_DRIVE_ID="drive-id",
    SHAREPOINT_FOLDER="Gastos/Validos",
)
class SharePointUploadTests(TestCase):
    @patch("validador.sharepoint.urlopen")
    @patch("validador.sharepoint.msal.ConfidentialClientApplication")
    def test_uploads_to_configured_drive_folder_without_overwriting(
        self, msal_application, urlopen
    ):
        confidential_client = MagicMock()
        confidential_client.acquire_token_for_client.return_value = {
            "access_token": "test-token"
        }
        msal_application.return_value = confidential_client

        graph_response = MagicMock()
        graph_response.__enter__.return_value.read.return_value = json.dumps(
            {
                "id": "item-id",
                "name": "gastos.csv",
                "webUrl": "https://sharepoint.example/gastos.csv",
            }
        ).encode()
        urlopen.return_value = graph_response

        result = upload_to_sharepoint(
            SimpleUploadedFile("gastos.csv", b"validated file"),
            "gastos.csv",
        )

        request = urlopen.call_args.args[0]
        self.assertIn("/drives/drive-id/root:/Gastos/Validos/gastos.csv:/content", request.full_url)
        self.assertIn("conflictBehavior=rename", request.full_url)
        self.assertEqual(request.data, b"validated file")
        self.assertEqual(request.get_header("Authorization"), "Bearer test-token")
        self.assertEqual(result["item_id"], "item-id")


    @override_settings(
        SHAREPOINT_TENANT_ID="tenant",
        SHAREPOINT_CLIENT_ID="client",
        SHAREPOINT_CLIENT_SECRET="secret",
        SHAREPOINT_DRIVE_ID="",
        SHAREPOINT_SITE_ID="site-id",
        SHAREPOINT_DOCUMENT_LIBRARY="Documentos",
    )
    @patch("validador.sharepoint.urlopen")
    @patch("validador.sharepoint.msal.ConfidentialClientApplication")
    def test_resolves_configured_document_library_from_site(
        self, msal_application, urlopen
    ):
        confidential_client = MagicMock()
        confidential_client.acquire_token_for_client.return_value = {
            "access_token": "test-token"
        }
        msal_application.return_value = confidential_client

        drives_response = MagicMock()
        drives_response.__enter__.return_value.read.return_value = json.dumps(
            {"value": [{"id": "documents-drive", "name": "Documentos"}]}
        ).encode()
        upload_response = MagicMock()
        upload_response.__enter__.return_value.read.return_value = json.dumps(
            {"id": "item-id", "name": "gastos.csv"}
        ).encode()
        urlopen.side_effect = [drives_response, upload_response]

        upload_to_sharepoint(
            SimpleUploadedFile("gastos.csv", b"validated file"),
            "gastos.csv",
        )

        self.assertIn(
            "/sites/site-id/drives",
            urlopen.call_args_list[0].args[0].full_url,
        )
        self.assertIn(
            "/drives/documents-drive/root:/Gastos/Validos/gastos.csv:/content",
            urlopen.call_args_list[1].args[0].full_url,
        )

    @override_settings(
        SHAREPOINT_TENANT_ID="tenant",
        SHAREPOINT_CLIENT_ID="client",
        SHAREPOINT_CLIENT_SECRET="secret",
        SHAREPOINT_DRIVE_ID="",
        SHAREPOINT_SITE_ID="site-id",
        SHAREPOINT_DOCUMENT_LIBRARY="",
    )
    def test_site_without_drive_or_library_name_is_not_configured(self):
        self.assertFalse(is_sharepoint_configured())


@override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
class MockLoginTests(TestCase):
    def test_protected_pages_redirect_to_login(self):
        response = self.client.get(reverse("index"))
        self.assertRedirects(response, f"{reverse('login')}?next=/")

    def test_login_page_shows_demo_username(self):
        response = self.client.get(reverse("login"))
        self.assertContains(response, "demo")
        self.assertContains(response, "En producción se desactiva")

    def test_valid_mock_credentials_create_session(self):
        response = self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        self.assertRedirects(response, reverse("design_preferences"))
        self.assertTrue(self.client.session.get("_auth_user_id"))

    def test_invalid_credentials_are_rejected(self):
        response = self.client.post(
            reverse("login"),
            {"username": "demo", "password": "wrong"},
        )
        self.assertEqual(response.status_code, 401)
        self.assertContains(response, "Usuario o contraseña incorrectos", status_code=401)

    @override_settings(DEBUG=False)
    def test_mock_login_is_disabled_when_debug_is_off(self):
        response = self.client.get(reverse("login"))
        self.assertEqual(response.status_code, 404)

    def test_logout_requires_post_and_clears_session(self):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )
        response = self.client.post(reverse("logout"))
        self.assertRedirects(response, reverse("login"))
        self.assertNotIn("_auth_user_id", self.client.session)


@override_settings(DEBUG=True, DEMO_LOGIN_ENABLED=True)
class DesignPreferenceTests(TestCase):
    def setUp(self):
        self.client.post(
            reverse("login"),
            {"username": "demo", "password": "Demo123!"},
        )

    def test_post_login_design_picker_displays_curated_themes(self):
        response = self.client.get(reverse("design_preferences"))
        self.assertContains(response, "Minimalista")
        self.assertContains(response, "Abstracto")
        self.assertContains(response, "Corporativo")

    def test_saving_theme_persists_and_redirects_to_app(self):
        response = self.client.post(
            reverse("design_preferences"),
            {"action": "save", "theme": "abstract"},
        )

        self.assertRedirects(
            response,
            reverse("dashboard"),
            fetch_redirect_response=False,
        )
        preference = UserDesignPreference.objects.get(user__username="demo")
        self.assertEqual(preference.theme, "abstract")
        app_response = self.client.get(reverse("index"))
        self.assertContains(app_response, 'data-theme="abstract"')

    def test_invalid_theme_is_not_saved(self):
        response = self.client.post(
            reverse("design_preferences"),
            {"action": "save", "theme": "custom-css"},
        )

        self.assertEqual(response.status_code, 400)
        self.assertFalse(
            UserDesignPreference.objects.filter(user__username="demo", theme="custom-css").exists()
        )

    @override_settings(AI_PROVIDER="ollama")
    @patch("validador.views.recommend_design", return_value="corporate")
    def test_ai_recommendation_saves_theme_and_redirects_to_applied_design(
        self, recommend_design
    ):
        response = self.client.post(
            reverse("design_preferences"),
            {
                "action": "recommend",
                "design_description": "Formal y sobrio para finanzas",
            },
        )

        recommend_design.assert_called_once_with("Formal y sobrio para finanzas")
        self.assertRedirects(
            response,
            reverse("dashboard"),
            fetch_redirect_response=False,
        )
        preference = UserDesignPreference.objects.get(user__username="demo")
        self.assertEqual(preference.theme, "corporate")

        app_response = self.client.get(reverse("index"))
        self.assertContains(app_response, 'data-theme="corporate"')

    @override_settings(OPENAI_API_KEY="test-key")
    @patch("validador.views.recommend_design")
    def test_ai_recommendation_requires_description_and_does_not_call_model(
        self, recommend_design
    ):
        response = self.client.post(
            reverse("design_preferences"),
            {"action": "recommend", "design_description": ""},
        )

        self.assertContains(response, "Describe brevemente", status_code=200)
        recommend_design.assert_not_called()
