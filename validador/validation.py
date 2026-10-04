import ast
import csv
import io
import re
import unicodedata
import operator
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

MAX_ROWS = 20_000
COSTA_RICAN_CURRENCY = "CRC"
REQUIRED_FIELDS = (
    "fecha",
    "proveedor",
    "numero_factura",
    "moneda",
    "subtotal",
    "impuesto",
    "total",
)
HEADER_ALIASES = {
    "date": "fecha",
    "fecha": "fecha",
    "provider": "proveedor",
    "vendor": "proveedor",
    "proveedor": "proveedor",
    "invoice_number": "numero_factura",
    "invoice_no": "numero_factura",
    "numero_factura": "numero_factura",
    "factura": "numero_factura",
    "currency": "moneda",
    "moneda": "moneda",
    "subtotal": "subtotal",
    "tax": "impuesto",
    "impuesto": "impuesto",
    "total": "total",
}
OUTPUT_FIELDS = ("estado", "errores")
EXCEL_FORMULA_OPERATORS = {
    ast.Add: operator.add,
    ast.Sub: operator.sub,
    ast.Mult: operator.mul,
    ast.Div: operator.truediv,
}
EXCEL_CELL_REFERENCE = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)", re.IGNORECASE)
EXCEL_SUM_RANGE = re.compile(
    r"SUM\(\s*(\$?[A-Z]{1,3}\$?\d+)\s*:\s*(\$?[A-Z]{1,3}\$?\d+)\s*\)",
    re.IGNORECASE,
)


class FileValidationError(ValueError):
    pass


def _normalize_header(value):
    text = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[\s-]+", "_", text)
    return HEADER_ALIASES.get(text, text)


def _unique_headers(values):
    headers = []
    seen = set()
    for index, value in enumerate(values, start=1):
        header = _normalize_header(value) or f"columna_{index}"
        unique_header = header
        suffix = 2
        while unique_header in seen:
            unique_header = f"{header}_{suffix}"
            suffix += 1
        headers.append(unique_header)
        seen.add(unique_header)
    return headers


def _read_rows(upload):
    suffix = upload.name.rsplit(".", 1)[-1].lower() if "." in upload.name else ""
    if suffix in ("csv", "tsv"):
        try:
            text = upload.read().decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise FileValidationError("El CSV debe estar codificado como UTF-8.") from exc
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel_tab if suffix == "tsv" else csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        first_row = next(reader, None)
        if not first_row:
            raise FileValidationError("El CSV está vacío o no tiene encabezados.")
        headers = _unique_headers(first_row)
        rows = [
            dict(zip(headers, row))
            for row in reader
            if any(value for value in row)
        ]
        return headers, rows

    if suffix == "xlsx":
        headers, rows, _ = _read_xlsx_table(upload)
        return headers, rows

    raise FileValidationError("Formato no admitido. Sube un archivo .csv, .tsv o .xlsx.")


def _read_xlsx_table(upload):
    tables = _read_xlsx_tables(upload)
    if not tables:
        raise FileValidationError("El Excel está vacío o no tiene encabezados.")
    best = max(
        tables,
        key=lambda table: (
            table["score"],
            table["month_count"],
            table["data_rows"],
            table["sheet_name"] == table["active_sheet"],
        ),
    )
    return best["headers"], best["rows"], best["info"]


def _read_xlsx_tables(upload):
    upload.seek(0)
    workbook = load_workbook(upload, read_only=False, data_only=False)
    try:
        candidates = []
        active_title = workbook.active.title
        for worksheet in workbook.worksheets:
            scan_rows = min(worksheet.max_row or 0, 50)
            if not scan_rows:
                continue
            preview = list(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=scan_rows,
                    values_only=True,
                )
            )
            for row_index, row in enumerate(preview):
                filled = [
                    (column_index, value)
                    for column_index, value in enumerate(row)
                    if value is not None and str(value).strip()
                ]
                if len(filled) < 2:
                    continue
                month_count = sum(
                    str(value).strip().casefold()
                    in {
                        "enero",
                        "febrero",
                        "marzo",
                        "abril",
                        "mayo",
                        "junio",
                        "julio",
                        "agosto",
                        "septiembre",
                        "octubre",
                        "noviembre",
                        "diciembre",
                    }
                    for _, value in filled
                )
                text_count = sum(isinstance(value, str) for _, value in filled)
                formula_count = sum(
                    isinstance(value, str) and value.startswith("=")
                    for _, value in filled
                )
                numeric_count = sum(
                    isinstance(value, (int, float))
                    for _, value in filled
                )
                semantic_header_count = sum(
                    _normalize_header(value)
                    in {
                        "mes",
                        "month",
                        "fecha",
                        "date",
                        "periodo",
                        "period",
                        "concepto",
                        "indicador",
                        "categoria",
                        "rubro",
                        "proveedor",
                        "total",
                    }
                    for _, value in filled
                )
                start_column = filled[0][0]
                end_column = filled[-1][0] + 1
                data_rows = sum(
                    any(value is not None for value in later[start_column:end_column])
                    for later in preview[row_index + 1 :]
                )
                score = (
                    len(filled) * 3
                    + text_count
                    + month_count * 5
                    + semantic_header_count * 10
                    - formula_count * 8
                    - numeric_count * 5
                )
                candidates.append(
                    (
                        score,
                        month_count,
                        data_rows,
                        worksheet.title == active_title,
                        worksheet,
                        row_index + 1,
                        start_column,
                        end_column,
                    )
                )

        if not candidates:
            raise FileValidationError("El Excel está vacío o no tiene encabezados.")

        tables = []
        for worksheet in workbook.worksheets:
            sheet_candidates = [
                candidate
                for candidate in candidates
                if candidate[4].title == worksheet.title
            ]
            if not sheet_candidates:
                continue
            (
                score,
                month_count,
                data_rows,
                _,
                _,
                header_row,
                start_column,
                end_column,
            ) = max(sheet_candidates, key=lambda candidate: candidate[:4])
            table_rows = worksheet.iter_rows(
                min_row=header_row,
                max_col=end_column,
                values_only=True,
            )
            header_values = next(table_rows, None)
            if not header_values:
                continue
            headers = _unique_headers(header_values[start_column:end_column])
            formula_cache = {}
            rows = []
            for row_number, row in enumerate(table_rows, start=header_row + 1):
                values = list(row[start_column:end_column])
                for offset, value in enumerate(values):
                    if isinstance(value, str) and value.startswith("="):
                        values[offset] = _evaluate_xlsx_formula(
                            value,
                            worksheet,
                            row_number,
                            start_column + offset + 1,
                            formula_cache,
                            set(),
                        )
                if any(value is not None for value in values):
                    rows.append(dict(zip(headers, values)))
            tables.append(
                {
                    "headers": headers,
                    "rows": rows,
                    "sheet_name": worksheet.title,
                    "active_sheet": active_title,
                    "score": score,
                    "month_count": month_count,
                    "data_rows": data_rows,
                    "info": {
                        "sheet_name": worksheet.title,
                        "header_row": header_row,
                    },
                }
            )
        return tables
    finally:
        workbook.close()


def _delimited_table_candidates(upload):
    suffix = upload.name.rsplit(".", 1)[-1].lower() if "." in upload.name else ""
    upload.seek(0)
    try:
        text = upload.read().decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise FileValidationError("El archivo delimitado debe estar codificado como UTF-8.") from exc
    if not text.strip():
        raise FileValidationError("El archivo está vacío.")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if suffix == "tsv" else csv.excel
    raw_rows = list(csv.reader(io.StringIO(text), dialect))
    nonempty_rows = [
        any(value is not None and str(value).strip() for value in values)
        for values in raw_rows
    ]
    candidates = []
    for header_index, row in enumerate(raw_rows[:50]):
        filled_headers = [
            value for value in row if value is not None and str(value).strip()
        ]
        if len(filled_headers) < 2:
            continue
        headers = _unique_headers(row)
        data_rows = raw_rows[header_index + 1 :]
        sample_limit = min(
            10,
            max(1, 10_000 // max(1, len(headers) * 120)),
        )
        sample_rows = [
            dict(zip(headers, values))
            for values in data_rows[:sample_limit]
            if any(value is not None and str(value).strip() for value in values)
        ]
        numeric_count = sum(isinstance(value, (int, float)) for value in row)
        score = len(filled_headers) * 3 + sum(
            isinstance(value, str) for value in filled_headers
        ) - numeric_count * 5
        candidates.append(
            {
                "headers": headers,
                "rows": sample_rows,
                "row_count": sum(nonempty_rows[header_index + 1 :]),
                "_raw_rows": raw_rows,
                "_header_index": header_index,
                "sheet_name": "CSV",
                "header_row": header_index + 1,
                "score": score,
            }
        )
    return candidates


def read_invoice_tables(upload):
    suffix = upload.name.rsplit(".", 1)[-1].lower() if "." in upload.name else ""
    if suffix in ("csv", "tsv"):
        candidates = _delimited_table_candidates(upload)
    elif suffix == "xlsx":
        candidates = [
            {
                "headers": table["headers"],
                "rows": table["rows"],
                "sheet_name": table["sheet_name"],
                "header_row": table["info"]["header_row"],
                "score": table["score"],
            }
            for table in _read_xlsx_tables(upload)
        ]
    else:
        raise FileValidationError(
            "Formato no admitido. Sube un archivo Excel (.xlsx), CSV o TSV."
        )
    if not candidates:
        raise FileValidationError(
            "No encontré una tabla con encabezados y filas. Revisa el archivo."
        )
    return candidates


def materialize_invoice_table_rows(table):
    raw_rows = table.get("_raw_rows")
    if raw_rows is None:
        return table["rows"]
    headers = table["headers"]
    header_index = table["_header_index"]
    return [
        dict(zip(headers, values))
        for values in raw_rows[header_index + 1 :]
        if any(value is not None and str(value).strip() for value in values)
    ]


def _evaluate_xlsx_formula(formula, worksheet, row, column, cache, visiting):
    coordinate = worksheet.cell(row=row, column=column).coordinate
    if coordinate in cache:
        return cache[coordinate]
    if coordinate in visiting or len(visiting) >= 64:
        return None
    visiting.add(coordinate)

    def cell_value(reference):
        match = EXCEL_CELL_REFERENCE.fullmatch(reference.replace("$", ""))
        if not match:
            raise ValueError("Referencia de celda no admitida.")
        column_letters, row_text = match.groups()
        ref_coordinate = f"{column_letters.upper()}{row_text}"
        ref_cell = worksheet[ref_coordinate]
        value = ref_cell.value
        if isinstance(value, str) and value.startswith("="):
            return _evaluate_xlsx_formula(
                value,
                worksheet,
                ref_cell.row,
                ref_cell.column,
                cache,
                visiting,
            )
        return _parse_decimal(value)

    try:
        expression = formula[1:]
        if len(expression) > 10_000:
            return None

        def sum_range(match):
            min_col, min_row, max_col, max_row = range_boundaries(
                f"{match.group(1).replace('$', '')}:{match.group(2).replace('$', '')}"
            )
            if (max_col - min_col + 1) * (max_row - min_row + 1) > 10_000:
                raise ValueError("El rango de suma supera el límite admitido.")
            total = 0
            for sum_row in range(min_row, max_row + 1):
                for sum_column in range(min_col, max_col + 1):
                    total += cell_value(
                        worksheet.cell(row=sum_row, column=sum_column).coordinate
                    ) or 0
            return str(total)

        expression = EXCEL_SUM_RANGE.sub(sum_range, expression)
        expression = EXCEL_CELL_REFERENCE.sub(
            lambda match: str(
                cell_value(f"{match.group(1)}{match.group(2)}") or 0
            ),
            expression,
        )
        if re.search(r"[^0-9eE.+\-*/()\s]", expression):
            return None
        parsed = ast.parse(expression, mode="eval")
        if sum(1 for _ in ast.walk(parsed)) > 256:
            return None

        def evaluate(node):
            if isinstance(node, ast.Expression):
                return evaluate(node.body)
            if isinstance(node, ast.Constant) and isinstance(
                node.value, (int, float)
            ):
                return Decimal(str(node.value))
            if isinstance(node, ast.UnaryOp) and isinstance(
                node.op, (ast.UAdd, ast.USub)
            ):
                value = evaluate(node.operand)
                return value if isinstance(node.op, ast.UAdd) else -value
            if isinstance(node, ast.BinOp) and type(node.op) in EXCEL_FORMULA_OPERATORS:
                return EXCEL_FORMULA_OPERATORS[type(node.op)](
                    evaluate(node.left), evaluate(node.right)
                )
            raise ValueError("La fórmula contiene operaciones no admitidas.")

        result = evaluate(parsed)
        cache[coordinate] = result if result.is_finite() else None
        return cache[coordinate]
    except (
        ArithmeticError,
        InvalidOperation,
        SyntaxError,
        ValueError,
        RecursionError,
    ):
        cache[coordinate] = None
        return None
    finally:
        visiting.discard(coordinate)


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for date_format in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            pass
    return None


def _parse_decimal(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None
    text = str(value).strip().replace(" ", "").replace("\u00a0", "")
    if not text:
        return None
    negative_parentheses = text.startswith("(") and text.endswith(")")
    if negative_parentheses:
        text = text[1:-1]
    text = re.sub(r"^[A-Za-z]{3}(?=[+-]?[\d,.])", "", text)
    text = re.sub(r"(?<=[\d,.])[A-Za-z]{3}$", "", text)
    text = re.sub(r"[$€£¥₡₱]", "", text)

    comma = text.rfind(",")
    period = text.rfind(".")
    if comma >= 0 and period >= 0:
        decimal_separator = "," if comma > period else "."
        grouping_separator = "." if decimal_separator == "," else ","
        text = text.replace(grouping_separator, "")
        if decimal_separator == ",":
            text = text.replace(",", ".")
    elif comma >= 0 or period >= 0:
        separator = "," if comma >= 0 else "."
        parts = text.split(separator)
        if len(parts) > 2 and all(len(part) == 3 for part in parts[1:]):
            text = "".join(parts)
        elif len(parts) > 2:
            text = "".join(parts[:-1]) + "." + parts[-1]
        elif len(parts) == 2 and len(parts[1]) == 3:
            text = "".join(parts)
        elif separator == ",":
            text = text.replace(",", ".")

    try:
        amount = Decimal(text)
        if negative_parentheses:
            amount = -amount
        return amount if amount.is_finite() else None
    except InvalidOperation:
        return None


def validate_upload(upload):
    headers, rows = _read_rows(upload)
    missing = [field for field in REQUIRED_FIELDS if field not in headers]
    if missing:
        raise FileValidationError(
            "Faltan columnas requeridas: " + ", ".join(missing)
        )
    if len(rows) > MAX_ROWS:
        raise FileValidationError(
            f"El archivo tiene más de {MAX_ROWS:,} filas; divídelo antes de procesarlo."
        )

    return validate_expense_rows(rows)


def validate_expense_rows(rows):
    seen_invoices = set()
    reported_rows = []
    issue_counts = {}

    for row_number, original in enumerate(rows, start=2):
        row = {
            field: original.get(field)
            for field in REQUIRED_FIELDS
        }
        errors = []

        if not str(row["proveedor"] or "").strip():
            errors.append("Falta proveedor")
        invoice_number = str(row["numero_factura"] or "").strip()
        if not invoice_number:
            errors.append("Falta número de factura")
        else:
            invoice_key = (
                str(row["proveedor"] or "").strip().casefold(),
                invoice_number.casefold(),
            )
            if invoice_key in seen_invoices:
                errors.append("Factura duplicada para el proveedor")
            seen_invoices.add(invoice_key)

        if _parse_date(row["fecha"]) is None:
            errors.append("Fecha inválida")
        currency = str(row["moneda"] or "").strip().upper()
        if currency != COSTA_RICAN_CURRENCY:
            errors.append("Moneda no admitida: solo se aceptan colones de Costa Rica (CRC)")

        amounts = {}
        for field in ("subtotal", "impuesto", "total"):
            amount = _parse_decimal(row[field])
            if amount is None:
                errors.append(f"{field.capitalize()} inválido")
            elif amount < 0:
                errors.append(f"{field.capitalize()} no puede ser negativo")
            else:
                amounts[field] = amount
        if len(amounts) == 3 and abs(
            amounts["subtotal"] + amounts["impuesto"] - amounts["total"]
        ) > Decimal("0.01"):
            errors.append("Total no coincide con subtotal + impuesto")

        for error in errors:
            issue_counts[error] = issue_counts.get(error, 0) + 1
        reported_rows.append(
            {
                **row,
                "fila": row_number,
                "estado": "Revisar" if errors else "Correcto",
                "errores": "; ".join(errors),
            }
        )

    return reported_rows, issue_counts


def build_csv_report(rows):
    output = io.StringIO(newline="")
    fields = ["fila", *REQUIRED_FIELDS, *OUTPUT_FIELDS]
    writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        safe_row = {
            field: (
                "'" + value
                if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r")
                else value
            )
            for field, value in row.items()
        }
        writer.writerow(safe_row)
    return "\ufeff" + output.getvalue()
